"""Generate the standalone Experiment 3 workbook and source-bound status notes.

Only reads the selected attempt and writes derived reports under this pilot's
results directory. Missing hardware results stay blank in Excel and null in
JSON. Native C++, Vitis C simulation, HLS, route and board are separate gates.
"""
from __future__ import annotations

import argparse
from copy import copy
from collections import Counter
from datetime import datetime, timedelta, timezone
import hashlib
import json
import math
from pathlib import Path
import re
import shutil
import struct
import sys
import xml.etree.ElementTree as ET

from tools.hardware_measurement_pilot.measurement import extrapolate, summarize_csv
from tools.hardware_measurement_pilot.schedule_scenarios import repeated_build_mix, schedule_pipeline


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "results/hardware_measurement_pilot_20261012"
DEFAULT_ATTEMPT = OUT / "attempt05"
PROTOCOL_VERSION = "D09_hardware_pilot_v4_post_pilot_draft_20261012"
ORIGINAL_WORKBOOK = Path("E:/1paper汇报/声呐_HWNAS_全链条分阶段梳理_v2_算子上板前置_20261010.xlsx")

D09 = [
    ("状态与用途", "DRAFT：草案可依据试做修订，并记录版本与差异；保留本次执行版、原始数据与条件偏差，不把CPU偏差追认为当时performance已达成。正式实验5冻结后不得倒改验收门槛。原工作簿F/G/H保持空白，M列只记试做说明；试做记录不进入正式成本表。", "原工作簿 先确定事项 M12/M14；用户最新指示；冻结runner及逐run manifest"),
    ("测试矩阵", "INT8的PW/DW3/DW5/nearest/add各2映射共10项；另加W8A16头部PW 2项、body8→head16重定标边界2项，共14项；每项5组夹具，共70组。W16权重推迟到D06确认，concat不在默认范围。", "用户最新范围；inventory.json；case.json"),
    ("真实层与精度", "头部PW：48×40×40→21×40×40；DW3/DW5：144×40×40；nearest：48×20×20→48×40×40；add：两路48×40×40；边界：48×40×40的I8→O16。DW/nearest/add只测INT8；头部混合档为I16/W8/O16。", "tools/modnas_elastic/cost_graph.py；quant_contracts.py"),
    ("整数规则", "有符号INT8/INT16、零点0；卷积权重均INT8；INT32累加按二进制补码模2^32回绕；INT64乘法重定标；nearest-even（半整数向偶数舍入）；最终饱和。add先分别对齐并饱和到公共INT8，再求和并重定标；混合边界显式读取INT8主体值并输出INT16。", "numeric_contract.json；reference.py；kernel.cpp"),
    ("布局与接口", "本试做核采用完整帧CHW连续布局、小端序；INT8为1字节、INT16为2字节，bias/params为小端INT32。通过AXI主接口访问DDR。该布局不能冒充历史pixel-major流接口的已测成本。", "case.json mapping.layout；kernel.cpp"),
    ("身份绑定", "保存case、全部源码、配置、输入/权重/bias/params、整数期望输出、numeric_contract、工具版本、器件、时钟、bitstream及hwh的SHA256；每条板上原始记录绑定case、numeric、bitstream、期望/实际输出哈希、实际PL时钟和board_id。", "source_manifest.json；measurement.CSV_FIELDS"),
    ("工具链与物理实现", "同一项使用一致版本的HLS与Vivado；当前目标xczu7ev-ffvc1156-2-e、2022.2。attempt05的PL配置187.5MHz、route周期16/3ns；HLS仍取更严格的5ns综合目标，两者分别记录。HLS核级估计、布线后core和含PS/AXI/计数器的集成route资源分列，不能互换。", "attempt05/board_packages/*/*/case/run_hls.tcl、build_shell.tcl、vivado_shell/shell_metadata.json"),
    ("仿真与比特流门槛", "Python整数参考→本地C++→Vitis C仿真均逐元素、逐字节一致；随后HLS、完整布线及setup/hold/pulse、约束覆盖、DRC均通过才接受bit。原始日志与失败尝试保留。", "native_summary.json；remote_evidence；shell_metadata.json"),
    ("相位与计算区间", "相位1=DDR加载；2=片上计算区间，含相位内局部控制开销；3=结果写回DDR；0=空闲/提交。当前证据为生成FSM与AXI adapter静态检查、HWH完整时钟/复位连接、独立timer RTL验证及同bit五夹具板端计数。完整动态集成RTL仿真为NOT_RUN。compute不能表述为仅算术指令周期，transfer不能表述为DDR总线占用。", "attempt05/board_packages/*/*/case/phase_validation.json、phase_evidence/{structure,review,probe_summary}.json；逐case绑定"),
    ("计数器有效性", "校验计数器ABI、序号严格递增、valid=1、busy=0，无周期溢出/相位错误/序号溢出；验证1→2→3→0时序和相位计数合计。超时、无进展、溢出均判本次失败，保留证据，禁止清除后重复到通过。", "attempt05/cases/*/board/选定run/counter_raw.jsonl、batch1..3/correctness_fixtures.json；phase_timer.v"),
    ("时钟与单位", "通过PYNQ Clocks读取PL0寄存器派生频率，记录每批及每组起止值并核验稳定性；周期按ns=cycles×10^9/f_PL换算，PL相位总时不得超过包含它的主机调用区间。此为寄存器频率证据和时间范围一致性检查，不声称独立物理频率校准（NOT_RUN）。目标频率或HLS预计周期不能替代实际读取与板上计数。", "shell_metadata.json；timing_manifest.json；Clocks.fclk0_mhz；independent_clock_scale_calibration"),
    ("三个计时范围", "compute_ns为相位2片上计算区间（含局部控制开销）；transfer_ns为相位1+3的加载/写回区间（含参数，不等同于DDR总线占用率）；call_ns为板上进程从调用开始到输出拷贝完成的单调时钟时间，包含参数配置、缓存维护、同步等待和输出拷贝。逐元素比较/hash、网络传输、日志I/O在call窗外。", "measurement.py；attempt05/cases/*/board/选定run/session.json"),
    ("夹具和板上正确性", "板上先逐一跑5组夹具：随机线性、随机激活、ties-even、溢出极值、shift0饱和；每组比较全部输出并核对二进制SHA256。延迟重复使用固定且显式绑定的一组输入，每次计时后在计时窗外比较全部输出；不能只查checksum/argmax。", "fixtures.txt；case.json fixtures；measurement.py"),
    ("独立批次与重复", "每项至少3个独立批次，每批重新加载bit或重启runtime并记录事件身份；每批至少50次预热、随后≥1000次正式调用。预热独立标记且不入统计；每批固定输入/缓冲区/时钟，同板不并发干扰。", "measurement.py independent_batches_minimum=3, warmup_minimum=50, measured_minimum=1000"),
    ("短算子分组计时", "短算子允许连续N次调用测总call时间/N；每次输出拷贝仍在组计时窗内，所有比较/hash在组计时窗外。保留实际N、总ns及每次PL周期。默认N=1；按50次预热host中位数<100us时选N=16，正式次数向上补齐整组（1008次=63组），N=1仍1000次。host每组均值只计一次，不复制N遍；若不同批次N不同，按N分层并标明混合口径。", "board runner；measurement.CSV_FIELDS含group_id/group_size/group_call_total_ns"),
    ("逐批环境记录", "记录实际PL0起止频率、CPU governor、起止温度与批次起止时间。有实际可写policy接口且支持时设performance，结束恢复原值；无接口/传感器时保留UNAVAILABLE或空值和条件偏差说明，不伪造状态，也不丢弃原始测量。", "manifest.batches；board environment records"),
    ("统计方法", "compute/transfer按每次PL观测、call按每组总时/N分别保留原始样本，报告均值、中位数、线性插值P95、样本SD(ddof=1)、CV。另报各批call中位数和相对批次中位数的最大偏差。不得去除慢样本或挑最快批次；NaN、非正值、缺失、重复序号、身份漂移和超时均保留并标记。", "measurement.statistics_ns；summarize"),
    ("异常派生标注", "原CSV不改，另存逐行注释：按case/batch/预热或正式/scope，以median+6×1.4826×MAD标注描述性高时延；host按唯一group总时/N只计一个观测再映射成员。MAD=0不作稳健异常判定，高于中位数只加非稳健描述标签。invalid行、全部慢样本及原顺序保留；标注不改变原统计、门槛或选样。", "anomaly_annotation.py；attempt05/anomaly_annotations/*/{timing_annotated.csv,annotation_rules.json}"),
    ("本次CPU条件偏差", "实际policy读到userspace、available仅userspace，记录实际频率。原冻结runner的UNAVAILABLE表示performance条件不可用，不能解释为policy未读取。数值/夹具/相位通过与protocol_conditions_met=false分别报告。草案可据此改为该板卡明确的userspace条件并记录实际频率/漂移门槛，或更换performance可用环境；正式实验5前冻结。此次原summary不追改为当时performance合规。", "attempt05/cases/*/board/选定run/batch*/environment_{start,end}.json及cpu_governor_setup.json"),
    ("通过条件与host解释", "所有夹具及全部正式样本输出100%逐位一致；每批正式调用≥1000、至少3批；PL compute/transfer逐批及合并CV≤5%。host以批间中位数稳定性和P95作描述性判断，不直接套5% CV，也不默认新增中位数硬阈值。RTL相位/实际时钟/接口证据和环境条件单列核验；统计PASS不等于全硬件链通过。", "measurement.py；板上夹具与计时仪表门槛"),
    ("阶段耗时", "每次尝试分别记录Python参考准备、本地C++编译/运行、Vitis csim、HLS综合、IP导出、Vivado综合、布局布线、门槛检查、bit生成、上板测试的实际墙钟秒；排队时间与共享准备时间单列。不能把本地C++运行时间填到板上latency。", "preparation.json；native_summary.json；hls_stages.jsonl；stage_elapsed.jsonl"),
    ("全量外推", "按真实物理实现去重后分层计数；按算子/精度/规模/映射记录完整阶段耗时。计入失败重试；单板串行、许可证/机器并行能力及每日可用小时明确说明。未覆盖类别或未完整上板时全量天数留空；理想并行下界不称日历预测。", "measurement.extrapolate；D07"),
    ("冻结与验证", "本次按DRAFT运行；14项、70组唯一夹具在各独立批次验证，逐批≥1000次计时；完整硬件链与RTL相位证据完成后复核，实验5前正式冻结。工具/内核/接口/精度/计时变更另立版本；不把试做范围登记写成正式冻结。", "DRAFT_EXECUTE_PILOT_FREEZE_BEFORE_EXPERIMENT5"),
]

D07 = [
    ("当前建议", "先以2种确实不同的低/高并行物理配置起步；若实测资源/延迟存在未覆盖的非支配折中，再增加第3种。依据是控制未闭合实现覆盖的测量规模，不把六个策略名当六次build。时间是否触发>21天，只以当前有数据的条件排程说明；不承诺65键全量完成日期。", "本次P4实测；extrapolation.json conditional_pilot_schedule"),
    ("策略名与物理配置", "reference_low/high/min_latency/min_dsp/min_lut/min_bram是从合格实测行中选择的6个策略名，可命中同一物理行。不能直接把6个名字乘到测量次数上。", "tools/fpga_h4/search.py choose；cost_api.py"),
    ("去重依据", "绑定并去重原语形状、输入/权重/输出位宽、数值/接口协议、PE、有效SIMD、tile、buffer、器件、工具版本、实际时钟和实现源码。元素算子的SIMD未参与运算时不能当独立物理维度。", "case.json mapping；fpga_h4/common.py MAPPING_KEY"),
    ("当前结构盘点", "本次启用W8A8和HEAD_W8A16试做域按结构/精度去重得65种原语；含D06待确认的HEAD_W16A16全登记域为69。每种确有2/3/6不同物理实现时，当前域名义数130/195/390；实际总数仍待绑定去重。", "用户最新范围；inventory.json；原工作簿D06"),
    ("覆盖缺口", "试做使用一个P4形状，不能自动代表所有空间尺寸。当前域仍包含maxpool、stem 3×3普通卷积以及不同尺寸/通道的原语；外推必须分层补齐或明确为局部估计。HEAD_W16A16若D06后续启用，需额外计数和独立测量。", "inventory.json counts_by_family_precision；D06"),
    ("当前整帧模板容量", "65个当前结构键中，18个仅显式BRAM输入/输出帧的存储下界就超过器件624 BRAM18，另4个缺模板（stem 3×3及maxpool）；剩余43个也只是未触发该下界，不能称已实现。例16×320×320→48×320×320的INT8 PW激活帧合计6.25MiB，至少2845 BRAM18。此限制针对当前整帧实现，不代表网络或层不能分块、流式或用URAM实现。", "buffer_coverage.json；buffer_coverage_independent_review.json；kernel.cpp BIND_STORAGE"),
    ("缩减触发", "仅当完整阶段实测、真实物理项数及可用算力/板时计划支持的全量日历时间超过21天，才考虑每种原语减少到2–3种实际映射。串行工作量超过21天只能触发复核，不能直接断言并行日历超过3周。", "用户实验3；measurement.extrapolate"),
    ("保留原则", "至少保留不同的低并行和高并行实现；请求PE翻倍不代表板上吞吐自动翻倍，应对照compute/transfer/host与core资源。若二者未覆盖资源/时延折中，第3个保留实测非支配点。六个查询策略仍可在较小的实测配置集合上选择。", "mapping_comparisons；reference_low/high板上数据及route core资源"),
    ("决策留痕", "记录缩减前后不同物理点、预计时间、资源/时序失败和覆盖损失；禁止以缺测配置的推测收益挑选映射。试做行与正式库分离，后续正式生产须重新按冻结协议采集。", "D09；独立决策记录（待实测）"),
]


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def read_json(path, default=None):
    path = Path(path)
    return json.loads(path.read_text(encoding="utf-8-sig")) if path.is_file() else default


def link(path, label=None):
    return f"[{label or Path(path).name}]({Path(path).resolve().as_posix()})"


def metadata_rows(rows):
    return "\n".join(f"| {key} | {text} | {source} |" for key, text, source in rows)


def _stages(path, case_id, source_label, issues):
    rows = []
    if not path.is_file():
        return rows
    for line, raw in enumerate(path.read_text(encoding="utf-8-sig").splitlines(), 1):
        try:
            record = json.loads(raw)
            seconds = record.get("wall_seconds", record.get("elapsed_seconds"))
            status = record.get("status") or ("PASS" if record.get("returncode") == 0 else "FAIL")
            rows.append(dict(case_id=case_id, stage=record.get("stage"), seconds=seconds,
                             status=status, evidence_layer=source_label, source=path.as_posix()))
        except (ValueError, TypeError) as error:
            issues.append(f"无法读取阶段记录 {path}:{line}: {error}")
    return rows


def _hls_resources(path):
    if not path.is_file():
        return {}
    tree = ET.parse(path)
    resources = tree.find(".//AreaEstimates/Resources")
    result = {}
    if resources is not None:
        for source, target in (("LUT", "LUT"), ("FF", "FF"), ("DSP", "DSP"), ("BRAM_18K", "BRAM18")):
            value = resources.findtext(source)
            result[target] = int(value) if value is not None else None
    return result


def _route_resources(path):
    if not path.is_file():
        return {}
    found = {}
    for line in path.read_text(encoding="utf-8", errors="replace").splitlines():
        fields = [item.strip() for item in line.strip().split("|")]
        if len(fields) < 4:
            continue
        name = fields[1].rstrip("*").strip()
        names = {"CLB LUTs": "LUT", "CLB Registers": "FF", "DSPs": "DSP", "Block RAM Tile": "BRAM36_equivalent"}
        if name in names and re.fullmatch(r"[\d,.]+", fields[2]):
            found.setdefault(names[name], float(fields[2].replace(",", "")))
    if "BRAM36_equivalent" in found:
        found["BRAM18"] = found["BRAM36_equivalent"] * 2
    return found


def _route_core_resources(path):
    """Read only an unambiguous core row from Vivado's hierarchy report.

    Column names, not fixed numeric positions, define resources. If a tool
    version emits a different table or several possible core rows, retain an
    empty result rather than guessing or subtracting shell overhead.
    """
    if not path.is_file():
        return {}
    headers, matches = None, []
    for line in path.read_text(encoding="utf-8", errors="replace").splitlines():
        cells = [value.strip() for value in line.strip().strip("|").split("|")]
        if "Instance" in cells and any("LUT" in value for value in cells):
            headers = cells
            continue
        if headers is None or len(cells) != len(headers):
            continue
        row = dict(zip(headers, cells))
        instance = row.get("Instance", "").strip()
        if instance.split("/")[-1] != "core":
            continue
        values = {}
        for column, value in row.items():
            if not re.fullmatch(r"[\d,.]+", value):
                continue
            number = float(value.replace(",", ""))
            if column in ("Total LUTs", "LUTs"):
                values["LUT"] = number
            elif column in ("FFs", "FF"):
                values["FF"] = number
            elif column.startswith("DSP"):
                values["DSP"] = number
            elif column in ("RAMB36", "RAMB18"):
                values[column] = number
        if "RAMB36" in values and "RAMB18" in values:
            values["BRAM18"] = 2 * values["RAMB36"] + values["RAMB18"]
        if all(name in values for name in ("LUT", "FF", "DSP", "BRAM18")):
            matches.append(dict(values, instance=instance))
    return matches[0] if len(matches) == 1 else {}


def _snapshot_root(path):
    """Find the immutable fetch root without searching unrelated board probes."""
    if path is not None:
        return next((parent for parent in Path(path).parents if (parent / "fetch_manifest.json").is_file()), None)
    return None


def _board_dispatches(attempt):
    result = {}
    for path in sorted((attempt / "board_transport").glob("*/transport.json")):
        transport = read_json(path, {})
        receipt = read_json(path.parent / "receipt_launch.json", {})
        if receipt.get("state") != "DISPATCHED":
            continue
        result.setdefault(transport.get("case_id"), []).append(dict(
            case_sha256=transport.get("case_sha256"), mode=transport.get("mode"),
            archive_sha256=transport.get("archive_sha256"), receipt=receipt,
            source=(path.parent / "receipt_launch.json").as_posix()))
    return result


def _reset_acceptance(bitstream_sha256, reset_audit):
    """A route PASS cannot override a separately verified reset defect."""
    matches = [row for row in reset_audit.get("cases", [])
               if bitstream_sha256 and row.get("bitstream_sha256") == bitstream_sha256]
    if any(row.get("board_acceptance") == "RESET_INTERFACE_INVALID_NOT_BOARD_ACCEPTED" for row in matches):
        return "RESET_INTERFACE_INVALID_NOT_BOARD_ACCEPTED"
    return "STATIC_INTERFACE_ACCEPTANCE_NOT_RECORDED"


def _package_case(path):
    if path is not None:
        return next((parent for parent in Path(path).parents
                     if parent.name == "case" and (parent / "case.json").is_file()
                     and (parent / "evidence").is_dir()), None)
    return None


def _bound_package_case(attempt, case, route_entry, bit_sha, source_manifest_sha):
    """Bind a local board package even when the newest tool record is a fetch."""
    if not route_entry or not bit_sha or not source_manifest_sha:
        return None, {}
    route = route_entry[1]
    expected = dict(case_id=case['case_id'], case_sha256=case['case_sha256'],
                    source_manifest_sha256=source_manifest_sha)
    if route.get('status') != 'PASS' or any(route.get(key) != value for key,value in expected.items()):
        return None, dict(status='NO_SOURCE_BOUND_ROUTE_RECORD')
    direct = _package_case(route_entry[0])
    candidates = ([direct] if direct else []) + sorted((attempt/'board_packages'/case['case_id']).glob('*/case'), reverse=True)
    seen, rejected = set(), []
    for package in candidates:
        if package.resolve() in seen:
            continue
        seen.add(package.resolve())
        transport_path = package/'package_transport_manifest.json'
        package_route_path = package/'evidence/route_status.json'
        transport,package_route = read_json(transport_path,{}),read_json(package_route_path,{})
        try:
            if (read_json(package/'case.json',{}) != case or
                    any(record.get(key) != value for record in (transport,package_route) for key,value in expected.items()) or
                    package_route.get('status') != 'PASS'):
                raise ValueError('package case/source/route identity differs')
            files=transport.get('files',{})
            if files.get('vivado_shell/top.bit') != bit_sha or files.get('evidence/route_status.json') != sha(package_route_path):
                raise ValueError('package bit or route-status manifest binding differs')
            for name,digest in files.items():
                target=(package/name).resolve()
                if not target.is_relative_to(package.resolve()) or not target.is_file() or sha(target)!=digest:
                    raise ValueError('package artifact hash differs: '+name)
            package_bit=next((row.get('sha256') for row in package_route.get('gate',{}).get('artifacts',[]) if row.get('path')=='vivado_shell/top.bit'),None)
            if package_bit != bit_sha:
                raise ValueError('package route bit differs from selected newest route bit')
            return package,dict(status='PASS_CASE_BIT_SOURCE_PACKAGE_BINDING', package=package.as_posix(),
                transport_manifest_sha256=sha(transport_path),package_route_status_sha256=sha(package_route_path),
                selected_route_status_source=route_entry[0].as_posix(),selected_route_status_sha256=sha(route_entry[0]),
                source_manifest_sha256=source_manifest_sha,bitstream_sha256=bit_sha,verified_transport_files=len(files))
        except (ValueError,KeyError,OSError) as error:
            rejected.append(dict(package=package.as_posix(),reason=str(error)))
    return None,dict(status='NO_MATCHING_CASE_BIT_SOURCE_BOUND_PACKAGE',rejected_candidates=rejected)


def _stage_directory(entry, fallback):
    if entry is None:
        return fallback
    path, record = entry
    return path.parent / record["stage"] if path.name in ("hls_status.json", "route_status.json") else path.parent


def _verified_probe(attempt, case, bit_sha, fixture_folder, issues):
    """Import correctness only from a hashed probe snapshot; never its counters."""
    matches = []
    for path in sorted((attempt / "board_snapshots").glob("*/evidence/output/probe_summary.json")):
        probe = read_json(path, {})
        if probe.get("case_sha256") != case.get("case_sha256") or probe.get("bitstream_sha256") != bit_sha:
            continue
        evidence = path.parent.parent
        snapshot = evidence.parent
        request = read_json(evidence / "request.json", {})
        if request.get("config", {}).get("mode") != "probe":
            continue
        try:
            manifest = read_json(evidence / "snapshot_manifest.json", {})
            download = read_json(snapshot / "download.json", {})
            if (manifest.get("schema") != "pilot_board_snapshot_v1" or
                    manifest.get("observed_state") != "PASS" or
                    download.get("archive_sha256") != sha(snapshot / "snapshot.tar.gz")):
                raise ValueError("probe snapshot/archive binding is missing")
            for name, digest in manifest.get("files", {}).items():
                item = (evidence / name).resolve()
                if not item.is_relative_to(evidence.resolve()) or not item.is_file() or sha(item) != digest:
                    raise ValueError("probe snapshot file hash differs: " + name)
            required = ["request.json", "output/probe_summary.json", "output/correctness_fixtures.json"]
            required += ["output/" + item["fixture_id"] + "_board.bin" for item in case["fixtures"]]
            if any(name not in manifest.get("files", {}) for name in required):
                raise ValueError("probe snapshot manifest omits required evidence")
            config = request.get("config", {})
            if (config.get("case_sha256") != case["case_sha256"] or
                    config.get("case_id") != case["case_id"] or
                    probe.get("status") != "PASS_BOARD_CORRECTNESS_PHASE_VALIDATION_PENDING" or
                    probe.get("latency_claim") != "NONE"):
                raise ValueError("probe request/result identity or scope differs")
            fixtures = read_json(path.parent / "correctness_fixtures.json", [])
            expected = {item["fixture_id"]: item for item in case["fixtures"]}
            if len(fixtures) != len(expected) or {item.get("fixture_id") for item in fixtures} != set(expected):
                raise ValueError("probe fixture coverage differs")
            for item in fixtures:
                name = item["fixture_id"]
                contract = expected[name]
                if (item.get("mismatched_elements") != 0 or
                        item.get("compared_elements") != contract["expected_elements"] or
                        item.get("observed_output_sha256") != contract["expected_output_sha256"] or
                        (path.parent / (name + "_board.bin")).read_bytes() != (fixture_folder / (name + "_expected.bin")).read_bytes()):
                    raise ValueError("probe output does not match complete integer fixture: " + name)
            matches.append(dict(status="PASS_PROBE", fixtures=len(fixtures),
                compared_elements=sum(item["compared_elements"] for item in fixtures), mismatched_elements=0,
                clock_hz=probe.get("clock_hz"), clock_hz_end=probe.get("clock_hz_end"),
                board_id=probe.get("board_id"), bitstream_sha256=bit_sha,
                summary_source=path.as_posix(), summary_sha256=sha(path),
                snapshot_manifest_source=(evidence / "snapshot_manifest.json").as_posix(),
                snapshot_manifest_sha256=sha(evidence / "snapshot_manifest.json"),
                created_utc=manifest.get("created_utc"), latency_claim="NONE", timing_status="NOT_RUN"))
        except (ValueError, KeyError, OSError) as error:
            issues.append(case["case_id"] + ": probe快照未接纳: " + str(error))
    return max(matches, key=lambda row: row.get("created_utc", "")) if matches else {}


def compare_replayed_summary(stored, replay):
    """Compare JSON data exactly, except finite computed statistics at <=8 ULP.

    ARM and Windows libm may round stdev/sqrt by one binary64 ULP. This is a
    representational comparison only: values are neither replaced nor rounded,
    metadata/thresholds stay exact, and every recomputed PASS/FAIL stays exact.
    """
    stored = json.loads(json.dumps(stored, allow_nan=False))
    replay = json.loads(json.dumps(replay, allow_nan=False))
    differences = []
    statistical_fields = {"mean_ns", "median_ns", "p95_ns", "sample_std_ns", "cv", "minimum_ns", "maximum_ns"}
    def computed_statistic(path):
        if len(path) < 4 or path[0] != "cases" or type(path[1]) is not int:
            return False
        tail = path[2:]
        if len(tail) >= 2 and tail[0] == "batches" and type(tail[1]) is int:
            tail = tail[2:]
        if len(tail) == 3 and tail[2] in statistical_fields:
            return ((tail[0] == "scopes" and tail[1] in ("compute", "transfer", "call")) or
                    (tail[0] == "host_scopes_by_group_size" and isinstance(tail[1], str) and tail[1].isdigit()))
        return (len(tail) == 2 and tail[0] == "host_median_stability" and tail[1] in
                ("median_of_batch_medians_ns", "maximum_relative_deviation")) or (
                len(tail) == 3 and tail[:2] == ("host_median_stability", "batch_medians_ns") and type(tail[2]) is int)
    def ordered_bits(value):
        bits = struct.unpack("!Q", struct.pack("!d", value))[0]
        sign = 1 << 63
        return sign - (bits & (sign - 1)) if bits & sign else sign + bits
    def check(actual, expected, path):
        label = "/".join(map(str, path)) or "<root>"
        if type(actual) is not type(expected):
            raise ValueError("summary replay type differs: " + label)
        if isinstance(expected, dict):
            if actual.keys() != expected.keys():
                raise ValueError("summary replay keys differ: " + label)
            for key in expected:
                check(actual[key], expected[key], path + (key,))
        elif isinstance(expected, list):
            if len(actual) != len(expected):
                raise ValueError("summary replay list length differs: " + label)
            for index, (left, right) in enumerate(zip(actual, expected)):
                check(left, right, path + (index,))
        elif isinstance(expected, float):
            if not math.isfinite(actual) or not math.isfinite(expected):
                raise ValueError("nonfinite summary float: " + label)
            if actual != expected:
                distance = abs(ordered_bits(actual) - ordered_bits(expected))
                if not computed_statistic(path) or distance > 8:
                    raise ValueError("summary replay float differs beyond allowed statistic ULP: " + label)
                differences.append(dict(path=label, stored=actual, replay=expected, ulp_distance=distance))
        elif actual != expected:
            raise ValueError("summary replay exact value differs: " + label)
    check(stored, replay, ())
    return dict(status="PASS_REPLAY_EQUIVALENCE", json_roundtrip_for_dictionary_keys=True,
        finite_statistical_float_max_ulp=8, comparison="IEEE754_BINARY64_ORDERED_INTEGER_DISTANCE",
        exact_fields="dictionary keys, list lengths/order, types, counts, identity, strings, booleans, metadata, thresholds, and all PASS/FAIL fields",
        tolerated_float_differences=differences, original_values_modified=False)


def _cpu_governor_display(output, batch_id):
    """Display actual policy reads separately from the frozen coarse summary."""
    details = []
    for name in ("environment_start.json", "environment_end.json"):
        environment = read_json(Path(output) / batch_id / name, {})
        cpu = environment.get("cpu", {})
        policies = cpu.get("policies", [])
        readable = [dict(path=item.get("path"), actual=item.get("actual"), available=item.get("available", []),
                         condition_deviation=item.get("condition_deviation")) for item in policies if item.get("actual")]
        details.append(dict(stage=name, raw_summary=cpu.get("cpu_governor"), policies=readable,
                            frequency_raw_khz=cpu.get("frequency_raw_khz", {})))
    observed = sorted({policy["actual"] for record in details for policy in record["policies"]})
    available = sorted({value for record in details for policy in record["policies"] for value in policy["available"]})
    performance_available = "performance" in available if observed else None
    return dict(observed_policy_values=observed, evidence=details,
        requested_governor="performance", available_policy_values=available,
        requested_governor_available=performance_available,
        display=("/".join(observed) + "（policy实读；performance条件不可用；原summary标记UNAVAILABLE）")
                if observed and any(record["raw_summary"] == "UNAVAILABLE" for record in details)
                else "/".join(observed) if observed else "UNAVAILABLE（未取得policy实际值）")


def validate_timing_snapshot(snapshot, package):
    """Read-only import audit. PASS means faithfully importable, not gate PASS."""
    from tools.hardware_measurement_pilot.board_runner import _case_identity, _phase_identity, _checked_files
    snapshot, package = Path(snapshot).resolve(), Path(package).resolve()
    evidence, output = snapshot / "evidence", snapshot / "evidence/output"
    download = read_json(snapshot / "download.json", {})
    manifest = read_json(evidence / "snapshot_manifest.json", {})
    if (download.get("archive_sha256") != sha(snapshot / "snapshot.tar.gz") or
            manifest.get("schema") != "pilot_board_snapshot_v1"):
        raise ValueError("timing snapshot archive/manifest binding differs")
    _checked_files(evidence, manifest.get("files"))
    needed = {"request.json", "output/summary.json", "output/session.json", "output/timing_manifest.json", "output/timing_raw.csv"}
    needed.update(f"output/batch{batch}/correctness_fixtures.json" for batch in (1, 2, 3))
    if not needed.issubset(manifest.get("files", {})):
        raise ValueError("timing snapshot is incomplete; leave unselected and preserve it")
    request = read_json(evidence / "request.json", {})
    case = _case_identity(package)
    if (request.get("config", {}).get("mode") != "timing" or
            request.get("config", {}).get("case_sha256") != case["case_sha256"]):
        raise ValueError("snapshot is not a timing request for this exact case")
    phase, _ = _phase_identity(package, case, package / "phase_validation.json")
    summary = read_json(output / "summary.json", {})
    timing = read_json(output / "timing_manifest.json", {})
    session = read_json(output / "session.json", {})
    case_specs = timing.get("cases", [])
    if len(case_specs) != 1:
        raise ValueError("one exact case is required per imported board run")
    expected = dict(case_sha256=case["case_sha256"], numeric_contract_sha256=case["numeric_contract_sha256"],
                    bitstream_sha256=phase["bitstream_sha256"])
    if any(case_specs[0].get(key) != value or session.get(key) != value for key, value in expected.items()):
        raise ValueError("timing manifest/session identity differs")
    if (session.get("phase_validation_sha256") != sha(package / "phase_validation.json") or
            session.get("board_bundle_sha256") != sha(package / "board_bundle.json") or
            summary.get("board_bundle_sha256") != sha(package / "board_bundle.json")):
        raise ValueError("timing phase/bundle identity differs")
    # Replay every derived statistic using the same frozen summarizer source.
    if sha(package / "measurement.py") != sha(ROOT / "tools/hardware_measurement_pilot/measurement.py"):
        raise ValueError("frozen measurement source differs; use its recorded version for replay")
    replay = summarize_csv(output / "timing_raw.csv", timing)
    replay.update(correctness_fixtures_passed=15, integrated_phase_validation="PASS_BOUND_ARTIFACT",
                  board_bundle_sha256=sha(package / "board_bundle.json"), protocol_status="DRAFT", formal_cost_eligible=False)
    comparison = compare_replayed_summary(summary, replay)
    expected_fixtures = {item["fixture_id"]: item for item in case["fixtures"]}
    for batch in (1, 2, 3):
        folder = output / f"batch{batch}"
        fixtures = read_json(folder / "correctness_fixtures.json", [])
        if len(fixtures) != len(expected_fixtures) or {row.get("fixture_id") for row in fixtures} != set(expected_fixtures):
            raise ValueError("timing correctness-fixture coverage differs")
        for row in fixtures:
            name = row["fixture_id"]
            spec = expected_fixtures[name]
            binary = f"output/batch{batch}/{name}_board.bin"
            if (binary not in manifest["files"] or row.get("mismatched_elements") != 0 or
                    row.get("compared_elements") != spec["expected_elements"] or
                    row.get("observed_output_sha256") != spec["expected_output_sha256"] or
                    (folder / (name + "_board.bin")).read_bytes() != (package / (name + "_expected.bin")).read_bytes()):
                raise ValueError("timing fixture does not match full integer output")
    return dict(status="PASS_IMPORT_AUDIT", case_id=case["case_id"], case_sha256=case["case_sha256"],
        snapshot=snapshot.as_posix(), output=output.as_posix(), snapshot_manifest_sha256=sha(evidence / "snapshot_manifest.json"),
        summary_sha256=sha(output / "summary.json"), raw_csv_sha256=sha(output / "timing_raw.csv"),
        numerical_status=summary.get("status"), protocol_conditions_met=summary.get("protocol_conditions_met"),
        hardware_measurement_eligible=summary.get("hardware_measurement_eligible"),
        condition_deviations=summary.get("condition_deviations", []),
        replay_comparison=comparison,
        cpu_governor_observations={f"batch{index}": _cpu_governor_display(output, f"batch{index}") for index in (1, 2, 3)},
        scope="READ_ONLY_SNAPSHOT_HASH_IDENTITY_FULL_FIXTURES_AND_RAW_STATISTICS_REPLAY")


def _pilot_schedule(cases, inventory, attempt):
    """Conditional P4 mix only: 3 build workers, one observed board runtime."""
    expected_ids=[row['case_id'] for row in inventory['pilot_cases']]
    if inventory.get('pilot_case_count') != 14 or len(expected_ids) != 14 or len(set(expected_ids)) != 14:
        raise ValueError('The conditional P4 mix requires the exact 14-case pilot inventory')
    case_index={row['case_id']:row for row in cases}
    if len(case_index) != len(cases) or set(case_index)-set(expected_ids):
        raise ValueError('Duplicate or unexpected case IDs cannot enter the conditional P4 mix')
    observations, missing = [], []
    for case_id in expected_ids:
        row=case_index.get(case_id)
        if row is None:
            missing.append(case_id)
            continue
        durations = [row.get(name) for name in ('hls_wall_seconds', 'route_wall_seconds', 'board_runtime_seconds', 'probe_wall_seconds')]
        valid_durations = all(isinstance(value, (int, float)) and not isinstance(value,bool) and math.isfinite(value) and value > 0 for value in durations)
        numeric_valid = (row['board_correctness'] == 'PASS_ALL_BATCH_FIXTURES' and row['board_timing'] == 'PASS' and
            row.get('phase_clock_evidence_verified') is True and row.get('phase_evidence',{}).get('status')=='PASS_BOUND_PHASE_EVIDENCE' and
            row.get('static_interface_acceptance')=='PASS_STATIC_ROUTE_RESET_IDENTITY' and
            row.get('board') in ('PASS','STATISTICS_PASS_WITH_CPU_GOVERNOR_DEVIATION'))
        if not valid_durations or not numeric_valid:
            missing.append(row['case_id'])
            continue
        observations.append(dict(case_id=row['case_id'], case_sha256=row['case_sha256'],
            build_seconds=durations[0]+durations[1], hls_seconds=durations[0], route_seconds=durations[1],
            board_seconds=durations[2]+durations[3], timing_worker_seconds=durations[2], probe_worker_seconds=durations[3],
            board_accepted=row.get('hardware_measurement_eligible') is True and row.get('board')=='PASS',
            board_observation_validated=True, condition_deviations=row.get('board_condition_deviations', []),
            source=dict(hls=row['hls_status_source'], route=row['route_status_source'],
                        board=row.get('board_runtime_source'), probe=row.get('probe_runtime_source'))))
    result = dict(status='READY_CONDITIONAL_OBSERVED_P4_MIX' if not missing else 'WAIT_COMPLETE_14_CASE_PILOT_MIX',
        full_measurement_status='NOT_ESTIMABLE', full_measurement_calendar_days=None,
        observed_case_count=len(observations), expected_pilot_case_count=len(expected_ids), missing_case_ids=missing,
        build_workers=3, board_workers=1, declared_build_workers=read_json(attempt/'parallel_launch.json', {}).get('jobs'),
        observations=observations, scenarios=[],
        board_scope='Observed probe worker + three-batch timing worker elapsed seconds, including their loading/correctness/measurement overhead; distinct from per-call latency',
        protocol_deviation_scope='Validated numeric/parity/phase observations retained with recorded userspace CPU deviation; not D09 condition compliance or formal-cost acceptance',
        assumptions=['Repeat the observed 14-case P4 low/high mix in fixed inventory order for nominal 65x2/3/6 items',
            'Use each observed HLS+route process duration in its recorded shared-machine context; these are not isolated serial benchmarks',
            'Three identical build workers; one board handles probe+timing after each build finishes',
            'No prediction for uncovered shapes, additional mapping behaviors, or unimplemented kernels',
            'Hours/8 and hours/24 are conditional available-work-day equivalents, not a promised calendar completion date'],
        excluded_time=['network upload/download and SSH/Jupyter startup outside the workers', 'manual phase review and idle queue gaps',
                       'new retries/failures, implementation redesign, board recovery and environmental remediation'],
        coverage_blockers=['18 current full-frame BRAM lower-bound violations', '4 primitives without a current template',
                           'exact physical mappings not yet bound/deduplicated', 'P4 mix not representative evidence for every shape'],
        mapping_recommendation='START_2_DISTINCT_PHYSICAL_CONFIGURATIONS_ADD_THIRD_IF_NONDOMINATED',
        full_measurement_time_trigger_over_21_days=None, conditional_time_trigger_over_21_workdays=None,
        formal_cost_eligible=False)
    if missing:
        return result
    for mappings in (2,3,6):
        count = inventory['unique_primitive_count'] * mappings
        jobs = []
        for index in range(count):
            source = observations[index % len(observations)]
            jobs.append(dict(job_id=str(index+1)+':'+source['case_id'], build_seconds=source['build_seconds'],
                board_seconds=source['board_seconds'], board_accepted=source['board_accepted'], board_observation_validated=True))
        schedule = schedule_pipeline(jobs, build_workers=3, allow_observed_protocol_deviations=True,
                                     implementation_coverage_complete=False, physical_bindings_complete=False)
        pipeline = schedule['conditional_pipeline_makespan_seconds']
        result['scenarios'].append(dict(distinct_physical_configurations_assumed=mappings, nominal_items=count,
            conditional_build_hours=schedule['conditional_build_makespan_seconds']/3600,
            single_board_runtime_hours=schedule['board_occupancy_seconds']/3600,
            conditional_pipeline_hours=pipeline/3600,
            conditional_workdays_at_8h=pipeline/(8*3600), conditional_workdays_at_24h=pipeline/(24*3600),
            full_measurement_calendar_days=None, excludes_unmeasured_overhead=True,
            schedule={key:value for key,value in schedule.items() if key!='jobs'}))
    result['conditional_time_trigger_over_21_workdays'] = any(row['conditional_workdays_at_8h'] > 21 for row in result['scenarios'])
    return result


def _mapping_comparisons(cases, timings, resources):
    timing_index = {(row['case_id'], row['scope']): row for row in timings}
    resource_index = {row['case_id']:row for row in resources}
    groups = {}
    for row in cases:
        groups.setdefault((row['operator'],row['input_bits'],row['weight_bits']), {})[row['mapping']] = row
    comparisons = []
    for (operator, input_bits, weight_bits), pair in sorted(groups.items(), key=lambda item:str(item[0])):
        low, high = pair.get('reference_low'), pair.get('reference_high')
        if not low or not high:
            continue
        item = dict(operator=operator, input_bits=input_bits, weight_bits=weight_bits,
            low_case=low['case_id'], high_case=high['case_id'],
            status='OBSERVED_PAIR' if low['board_timing']==high['board_timing']=='PASS' else 'PAIR_NOT_COMPLETE',
            scope='Same-operator observed low/high pilot comparison, not a network benefit or automatic PE scaling claim')
        for name, case in (('low',low),('high',high)):
            item[name] = dict(compute_mean_ns=timing_index[(case['case_id'],'compute')]['mean_ns'],
                transfer_mean_ns=timing_index[(case['case_id'],'transfer')]['mean_ns'],
                host_median_ns=timing_index[(case['case_id'],'call')]['median_ns'],
                core_resources=resource_index[case['case_id']]['route_core'], board_status=case['board'])
        for scope, field in (('compute','compute_mean_ns'),('transfer','transfer_mean_ns'),('host_median','host_median_ns')):
            left,right = item['low'][field],item['high'][field]
            item[scope+'_low_over_high'] = left/right if item['status']=='OBSERVED_PAIR' and left and right else None
        item['observed_compute_high_minus_low_percent'] = ((item['high']['compute_mean_ns']/item['low']['compute_mean_ns']-1)*100
            if item['status']=='OBSERVED_PAIR' and item['low']['compute_mean_ns'] else None)
        comparisons.append(item)
    return comparisons


def _phase_evidence_info(attempt, case, bit_sha, issues):
    from tools.hardware_measurement_pilot.board_runner import _phase_identity
    for path in sorted((attempt/'board_packages'/case['case_id']).glob('*/case/phase_validation.json'), reverse=True):
        phase = read_json(path,{})
        if phase.get('case_sha256') != case.get('case_sha256') or phase.get('bitstream_sha256') != bit_sha:
            continue
        try:
            checked, files = _phase_identity(path.parent,case,path)
            return dict(status='PASS_BOUND_PHASE_EVIDENCE', source=path.as_posix(), source_sha256=sha(path),
                bound_files=len(files), method=checked.get('method'), admitted_scope=checked.get('admitted_scope'),
                dynamic_integrated_rtl_simulation=checked.get('dynamic_integrated_rtl_simulation'),
                excluded_claims=checked.get('excluded_claims'), actual_probe_clock_hz=checked.get('actual_probe_clock_hz'))
        except (ValueError,KeyError,OSError) as error:
            issues.append(case['case_id']+': phase validation来源绑定未通过: '+str(error))
    return dict(status='NO_BOUND_PHASE_EVIDENCE_RECORDED')


def _annotation_index(attempt, cases, issues):
    path = attempt/'anomaly_annotations/selected_runs_index.json'
    index = read_json(path,{})
    case_index = {row['case_id']:row for row in cases}
    checked = []
    for row in index.get('records',[]):
        current = case_index.get(row.get('case_id'))
        rules_path,csv_path = Path(row['rules']),Path(row['annotated_csv'])
        rules = read_json(rules_path,{})
        if (not current or row.get('raw_sha256') != current.get('board_raw_csv_sha256') or
                rules.get('input_sha256') != row.get('raw_sha256') or not csv_path.is_file() or
                rules.get('output_sha256') != sha(csv_path) or row.get('input_rows') != row.get('output_rows')):
            issues.append(str(row.get('case_id'))+': 异常派生标注与当前选定run的哈希/行数绑定未通过')
            continue
        checked.append(dict(row,flag_row_counts=rules.get('flag_row_counts',{}),
            rule_status='BOUND_DESCRIPTIVE_ANNOTATIONS',rules_sha256=sha(rules_path),output_sha256=sha(csv_path)))
    return dict(source=path.as_posix(), source_sha256=sha(path) if path.is_file() else None,
        records=checked,case_count=len(checked), total_rows=sum(row['input_rows'] for row in checked),
        all_raw_rows_preserved=index.get('all_raw_rows_preserved'),statistical_exclusions=index.get('statistical_exclusions'),
        scope='Descriptive labels only; MAD-zero above-median labels are not robust outlier decisions')


def _number(value, digits=4, divisor=1):
    return '—' if value is None else f'{value/divisor:.{digits}f}'


def _schedule_markdown(data):
    plan=data['extrapolation']['conditional_pilot_schedule']
    if not plan['scenarios']:
        return f"当前已有{plan['observed_case_count']}/{plan['expected_pilot_case_count']}项完整构建＋probe＋三批timing占用记录；完整P4试做mix尚未齐全，条件排程数值暂空。"
    rows=[]
    for row in plan['scenarios']:
        rows.append('| '+' | '.join([str(row['nominal_items']),*[_number(row[name],3) for name in
            ('conditional_build_hours','single_board_runtime_hours','conditional_pipeline_hours','conditional_workdays_at_8h','conditional_workdays_at_24h')]])+' |')
    return '\n'.join(['| 名义物理项数 | 三并行构建 h | 单板累计占用 h | 构建→单板排程 h | 已测工具+板端worker阶段条件工作日（8h/日） | 已测工具+板端worker阶段条件工作日（24h/日） |',
        '|---|---|---|---|---|---|',*rows])


def _comparison_markdown(data):
    rows=[]
    for item in data['mapping_comparisons']:
        label=f"{item['operator']} I{item['input_bits']}/W{item['weight_bits'] or '-'}"
        rows.append('| '+' | '.join([label,
            _number(item['low']['compute_mean_ns'],6,1e6),_number(item['high']['compute_mean_ns'],6,1e6),
            _number(item['compute_low_over_high'],6),_number(item['low']['host_median_ns'],6,1e6),
            _number(item['high']['host_median_ns'],6,1e6),item['status']])+' |')
    return '\n'.join(['| 同算子精度 | low compute ms | high compute ms | low/high | low host中位数 ms | high host中位数 ms | 状态 |',
        '|---|---|---|---|---|---|---|',*rows])


def _d07_rows(data):
    plan=data['extrapolation']['conditional_pilot_schedule']
    if plan['scenarios']:
        max_days=max(row['conditional_workdays_at_8h'] for row in plan['scenarios'])
        statement=(f"14项P4已观测mix按3个构建worker和单板串行占用模拟，最大名义390项的已测工具+板端worker阶段条件工作日（8h/日）为{max_days:.3f}日。"
            + ('仅此已测阶段条件情景未触发>21日削减证据；不能据此断言完整流程不超过3周。' if not plan['conditional_time_trigger_over_21_workdays'] else
               '仅此已测阶段条件情景触发>21工作日复核，须再核对真实物理去重与完整日历。')
            + 'userspace偏差、覆盖18个BRAM下界超限/4个缺模板和桥接、人工复核、修复等未计开销仍保留；完整全量可信日期为NOT_ESTIMABLE。')
    else:
        statement=f"已有{plan['observed_case_count']}/14项完整P4过程占用，等待其余真实结果后计算条件排程；当前不宣称已触发>21天。完整全量可信日期仍为NOT_ESTIMABLE。"
    return [D07[0],('实测时间条件',statement,'extrapolation.json conditional_pilot_schedule；BOARD_WORKER及HLS/route原始秒数'),*D07[1:]]


def _driver_observations(attempt, cases, issues):
    """Import completed automatic-driver envelopes; exclude the manual first case."""
    case_index = {row['case_id']:row for row in cases}
    records, pending = [], []
    for event_path in sorted((attempt/'board_execution').glob('*/events.jsonl')):
        result_path = event_path.parent/'result.json'
        inputs_path = event_path.parent/'inputs.json'
        result, inputs = read_json(result_path,{}), read_json(inputs_path,{})
        if not result:
            pending.append(event_path.parent.name)
            continue
        case = case_index.get(result.get('case_id'))
        if (not case or result.get('status') != 'PASS_IMPORT_AUDIT' or
                result.get('case_sha256') != case['case_sha256'] or
                inputs.get('case_sha256') != case['case_sha256'] or
                result.get('raw_csv_sha256') != case.get('board_raw_csv_sha256') or
                result.get('summary_sha256') != case.get('board_summary_sha256')):
            issues.append(event_path.parent.name+': 自动driver结果未绑定当前选定run，未纳入耗时描述')
            continue
        try:
            events = [json.loads(line) for line in event_path.read_text(encoding='utf-8').splitlines() if line.strip()]
            times = [datetime.fromisoformat(event['utc']) for event in events]
            if (not events or any(event.get('case_id') != case['case_id'] for event in events) or
                    any(time.tzinfo is None for time in times) or times != sorted(times)):
                raise ValueError('case身份、时区或事件时间顺序异常')
            def unique(stage):
                hits = [time for event,time in zip(events,times) if event.get('stage') == stage]
                if len(hits) != 1:
                    raise ValueError(stage+'事件不唯一')
                return hits[0]
            start,finish = unique('probe_upload_begin'),unique('timing_imported')
            elapsed = (finish-start).total_seconds()
            worker = case.get('board_runtime_seconds'),case.get('probe_wall_seconds')
            if not all(isinstance(value,(int,float)) and not isinstance(value,bool) and math.isfinite(value) and value>0 for value in worker):
                raise ValueError('缺少完整probe与timing worker耗时')
            outside = elapsed-sum(worker)
            if elapsed<=0 or outside<0:
                raise ValueError('driver包络小于其worker合计耗时')
            records.append(dict(case_id=case['case_id'],case_sha256=case['case_sha256'],
                driver_run=event_path.parent.name, started_utc=start.isoformat(),finished_utc=finish.isoformat(),
                driver_elapsed_seconds=elapsed, probe_plus_timing_worker_seconds=sum(worker),
                outside_worker_residual_seconds=outside,
                source=event_path.as_posix(), source_sha256=sha(event_path),
                import_result_source=result_path.as_posix(), import_result_sha256=sha(result_path),
                inputs_source=inputs_path.as_posix(),inputs_sha256=sha(inputs_path)))
        except (ValueError,KeyError,TypeError) as error:
            issues.append(event_path.parent.name+': 自动driver耗时未纳入: '+str(error))
    def descriptive(field):
        values=[row[field] for row in records]
        return dict(n=len(values),mean_seconds=sum(values)/len(values) if values else None,
                    minimum_seconds=min(values) if values else None,maximum_seconds=max(values) if values else None)
    return dict(status='OBSERVED_COMPLETED_AUTOMATIC_DRIVER_ENVELOPES',case_count=len(records),records=records,
        incomplete_driver_runs=pending,driver_elapsed=descriptive('driver_elapsed_seconds'),
        outside_worker_residual=descriptive('outside_worker_residual_seconds'),
        scope='Actual probe-upload start through timing-import finish; excludes manual first body8-high case and incomplete runs',
        residual_scope='Driver envelope minus recorded probe/timing workers: bridge, polling, local validation and orchestration combined; not pure network overhead',
        extrapolated_to_all_14=False,used_in_conditional_schedule=False)


def collect(attempt):
    issues = []
    inventory = read_json(attempt / "inventory.json") or read_json(OUT / "inventory.json")
    if not inventory:
        raise ValueError("prepared inventory.json is required")
    if inventory.get("schema_version") != "hardware_measurement_pilot_inventory_v2" or inventory.get("pilot_case_count") != 14:
        raise ValueError("Current report requires the revised 14-case W8A8/W8A16 inventory; preserve attempt01 as superseded")
    source_manifest = attempt / "source_manifest.json"
    source_manifest_sha = sha(source_manifest) if source_manifest.is_file() else None
    native = read_json(attempt / "native_summary.json", {})
    native_bound = source_manifest.is_file() and native.get("source_manifest_sha256") == sha(source_manifest)
    if native and not native_bound:
        issues.append("本地C++汇总未绑定当前source_manifest；旧PASS不提升为当前PASS。")
    native_cases = {row["case_id"]: row for row in native.get("cases", [])}
    remote = defaultdict_status(attempt / "remote_evidence")
    reset_audit_path = OUT / "diagnostics/aux_reset_failure_20261012/reset_interface_audit.json"
    reset_audit = read_json(reset_audit_path, {})
    dispatches = _board_dispatches(attempt)
    unreachable = read_json(OUT / "preflight/board_unreachable_diagnostic.json", {})
    archived_unreachable = bool(unreachable.get("reachability")) and not any(
        row.get("status") in ("PASS", "CONNECTED", "OPEN") for row in unreachable.get("reachability", []))
    stage_rows, cases, timing_rows, resources, batch_rows, host_strata = [], [], [], [], [], []
    for planned in inventory["pilot_cases"]:
        case_id = planned["case_id"]
        folder = attempt / "cases" / case_id
        case = read_json(folder / "case.json", planned)
        record = native_cases.get(case_id, {})
        native_state = record.get("status", "NOT_RUN")
        if native_state == "PASS" and (not native_bound or record.get("case_sha256") != case.get("case_sha256")):
            native_state = "STALE_SOURCE_BINDING"
        hls_entry = remote.get((case_id, "hls"))
        route_entry = remote.get((case_id, "route"))
        stages = {}
        for name, entry in (("hls", hls_entry), ("route", route_entry)):
            if not entry:
                stages[name] = "NOT_RUN"
            elif (entry[1].get("case_sha256") != case.get("case_sha256") or
                  entry[1].get('source_manifest_sha256') != source_manifest_sha):
                stages[name] = "STALE_SOURCE_BINDING"
            else:
                stages[name] = entry[1].get("status", "UNKNOWN")
        hls_dir = _stage_directory(hls_entry, folder)
        route_dir = _stage_directory(route_entry, folder / "vivado_shell")
        csim = read_json(hls_dir / "csim_parity.json", {})
        csim_state = csim.get("status", "NOT_RUN")
        if stages["hls"] == "STALE_SOURCE_BINDING":
            csim_state = "STALE_SOURCE_BINDING"
        shell_candidates = ([route_dir / "shell_metadata.json"] + sorted(route_dir.rglob("shell_metadata.json"))
                            if route_dir.exists() else []) + [folder / "vivado_shell" / "shell_metadata.json"]
        shell_path = next((path for path in shell_candidates if path.is_file()), shell_candidates[0])
        shell = read_json(shell_path, {})
        snapshot = _snapshot_root(route_entry[0] if route_entry else None)
        remote_bit = next((row for row in (route_entry[1].get("gate", {}).get("artifacts", []) if route_entry else [])
                           if row.get("path") == "vivado_shell/top.bit"), {})
        bit_sha = remote_bit.get("sha256") if stages["route"] == "PASS" else None
        package,package_binding = _bound_package_case(attempt,case,route_entry,bit_sha,source_manifest_sha)
        bit_candidates = ([route_dir / "top.bit"] + sorted(route_dir.rglob("top.bit"))
                          if route_dir.exists() else []) + [folder / "vivado_shell" / "top.bit"]
        if snapshot:
            bit_candidates.append(snapshot / "cases" / case_id / "vivado_shell" / "top.bit")
        if package:
            bit_candidates.append(package / "vivado_shell" / "top.bit")
        bit = next((path for path in bit_candidates if path.is_file()), None)
        bit_state = "PRESENT_UNVALIDATED" if bit else "NOT_GENERATED"
        if bit_sha and shell.get("route_timing_drc") == "PASS":
            bit_state = "GENERATED_REMOTE_SHA_RECORDED"
            if bit:
                bit_state = "GENERATED_LOCAL_SHA_VERIFIED" if sha(bit) == bit_sha else "BITSTREAM_SHA_MISMATCH"
                if bit_state == "BITSTREAM_SHA_MISMATCH":
                    issues.append(case_id + ": 本地bit与route记录SHA不匹配。")
        interface_acceptance = _reset_acceptance(bit_sha, reset_audit)
        interface_evidence = {}
        if (interface_acceptance == "STATIC_INTERFACE_ACCEPTANCE_NOT_RECORDED" and
                bit_state == "GENERATED_LOCAL_SHA_VERIFIED" and package is not None):
            from tools.hardware_measurement_pilot.board_runner import _case_identity, _route_identity, _reset_wiring_identity
            try:
                _case_identity(package)
                _route_identity(package)
                interface_evidence = _reset_wiring_identity(package / "vivado_shell" / "top.hwh")
                interface_evidence.update(validator_sha256=sha(ROOT / "tools/hardware_measurement_pilot/board_runner.py"),
                    hwh_source=(package / "vivado_shell" / "top.hwh").as_posix(),
                    package_binding=package_binding,
                    scope="LOCAL_STATIC_ROUTE_RESET_IDENTITY_ONLY_NO_BOARD_EXECUTION")
                phase_path = package.parent / "phase_structure.json"
                phase = read_json(phase_path, {})
                if (phase.get("case_sha256") == case.get("case_sha256") and
                        phase.get("wiring", {}).get("hwh_sha256") == interface_evidence["hwh_sha256"]):
                    interface_evidence["phase_structure_review"] = dict(source=phase_path.as_posix(),
                        source_sha256=sha(phase_path), status=phase.get("status"),
                        clock_port_count=len(phase.get("wiring", {}).get("clock_ports", [])),
                        single_clock_domain=phase.get("wiring", {}).get("single_clock_domain"),
                        reset_status=phase.get("wiring", {}).get("reset_wiring", {}).get("status"),
                        excluded_claims=phase.get("excluded_claims", []))
                interface_acceptance = "PASS_STATIC_ROUTE_RESET_IDENTITY"
            except (ValueError, KeyError, ET.ParseError) as error:
                interface_acceptance = "STATIC_ROUTE_RESET_IDENTITY_FAILED"
                interface_evidence = dict(error=str(error))
                issues.append(case_id + ": 静态route/reset核验失败: " + str(error))
        board_folder = folder / "board"
        selected_run = read_json(board_folder / "selected_run.json", {})
        if selected_run:
            selected_folder = (board_folder / selected_run.get("run_directory", "")).resolve()
            if (selected_run.get("evidence_kind") != "BOARD_MEASUREMENT" or
                    not selected_folder.is_relative_to(board_folder.resolve()) or selected_folder == board_folder.resolve()):
                raise ValueError("Invalid explicit board measurement run pointer: " + case_id)
            board_folder = selected_folder
        board_summary_path = board_folder / "summary.json"
        board_summary = read_json(board_summary_path, {})
        board_case = next((r for r in board_summary.get("cases", []) if r.get("case_id") == case_id), {})
        board_status = board_case.get("status", "NOT_RUN")
        fixture_batches = [read_json(board_folder / f"batch{index}" / "correctness_fixtures.json", []) for index in (1, 2, 3)]
        fixture_records = [item for batch in fixture_batches for item in batch]
        fixture_expected = {item["fixture_id"]: item for item in case.get("fixtures", [])}
        fixture_ok = bool(fixture_expected) and all(len(batch) == len(fixture_expected) and
            {item.get("fixture_id") for item in batch} == set(fixture_expected) for batch in fixture_batches) and all(
            item.get("fixture_id") in fixture_expected and item.get("mismatched_elements") == 0 and
            item.get("observed_output_sha256") == fixture_expected[item["fixture_id"]]["expected_output_sha256"]
            for item in fixture_records)
        board_manifest = read_json(board_folder / "timing_manifest.json", {})
        board_spec = next((r for r in board_manifest.get("cases", []) if r.get("case_id") == case_id), {})
        raw = board_folder / "timing_raw.csv"
        board_bound = (board_spec.get("case_sha256") == case.get("case_sha256") and raw.is_file() and
                       board_summary.get("raw_csv_sha256") == sha(raw))
        if board_summary and not board_bound:
            board_status = "STALE_SOURCE_BINDING"
            issues.append(case_id + ": 板上summary的case/raw CSV绑定缺失或不匹配。")
        # Statistics are separate from the full board/phase/environment gate.
        if board_status == "PASS" and not (fixture_ok and board_summary.get("status") == "PASS" and
                                          board_summary.get("hardware_measurement_eligible") is True):
            deviations = board_summary.get("condition_deviations", [])
            only_cpu = bool(deviations) and all(row.get("condition") == "cpu_governor" or
                (row.get("condition") == "runner_recorded_condition" and any(marker in row.get("observed", "")
                    for marker in ("NOT_PERFORMANCE", "PERFORMANCE_GOVERNOR_NOT_AVAILABLE", "CPU_GOVERNOR"))) for row in deviations)
            if not fixture_ok:
                board_status = "STATISTICS_PASS_BATCH_FIXTURES_PENDING"
            elif board_summary.get("status") != "PASS":
                board_status = "STATISTICS_PASS_SUMMARY_GATE_FAILED"
            elif board_summary.get("phase_clock_evidence_verified") is not True:
                board_status = "STATISTICS_PASS_PHASE_CLOCK_GATE_PENDING"
            elif only_cpu:
                board_status = "STATISTICS_PASS_WITH_CPU_GOVERNOR_DEVIATION"
            elif deviations:
                board_status = "STATISTICS_PASS_WITH_ENVIRONMENT_DEVIATIONS"
            else:
                board_status = "STATISTICS_PASS_HARDWARE_ELIGIBILITY_NOT_RECORDED"
        probe = _verified_probe(attempt, case, bit_sha, folder, issues)
        phase_evidence = _phase_evidence_info(attempt, case, bit_sha, issues)
        if board_status.startswith('STATISTICS_PASS') or board_status=='PASS':
            if phase_evidence['status'] != 'PASS_BOUND_PHASE_EVIDENCE':
                board_status = 'STATISTICS_PASS_PHASE_BINDING_PENDING'
        board_runtime_path = board_folder/'board_stages.jsonl'
        runtime_rows = _stages(board_runtime_path,case_id,'BOARD_WORKER',issues) if board_bound else []
        probe_runtime_path = Path(probe['summary_source']).parent/'board_stages.jsonl' if probe else None
        probe_rows = _stages(probe_runtime_path,case_id,'BOARD_PROBE_WORKER',issues) if probe_runtime_path else []
        step_rows_for_board = runtime_rows + probe_rows
        stage_rows.extend(step_rows_for_board)
        runtime_seconds = next((row['seconds'] for row in runtime_rows if row['stage']=='board_test' and row['status']=='PASS'),None)
        probe_seconds = next((row['seconds'] for row in probe_rows if row['stage']=='board_probe' and row['status']=='PASS'),None)
        bound_dispatches = [row for row in dispatches.get(case_id, []) if row["case_sha256"] == case.get("case_sha256")]
        execution_state = "RESULT_RECORDED" if board_summary else "NOT_DISPATCHED"
        if bound_dispatches and not board_summary:
            # The archived outage followed the rejected reset-invalid package.
            # It is not evidence of reachability after a new reset-corrected
            # attempt, even when both attempts use the same logical case_id.
            execution_state = ("DISPATCHED_NO_RESULT_UNREACHABLE_IN_ARCHIVED_DIAGNOSTIC" if archived_unreachable
                               and interface_acceptance == "RESET_INTERFACE_INVALID_NOT_BOARD_ACCEPTED"
                               else "DISPATCHED_NO_RESULT")
        if probe and not board_summary:
            execution_state = "PROBE_RESULT_RECORDED"
        board_acceptance = (interface_acceptance if interface_acceptance == "RESET_INTERFACE_INVALID_NOT_BOARD_ACCEPTED"
                            else "PASS" if board_status == "PASS" and interface_acceptance == "PASS_STATIC_ROUTE_RESET_IDENTITY"
                            else "VALIDATED_RUNTIME_WITH_CPU_GOVERNOR_DEVIATION" if board_status == "STATISTICS_PASS_WITH_CPU_GOVERNOR_DEVIATION"
                                 and interface_acceptance == "PASS_STATIC_ROUTE_RESET_IDENTITY"
                            else "RUNTIME_PASS_STATIC_INTERFACE_ACCEPTANCE_NOT_RECORDED" if board_status == "PASS"
                            else "NOT_BOARD_ACCEPTED")
        cases.append(dict(case_id=case_id, operator=case["operator"], input_bits=case["input_bits"],
                          weight_bits=case.get("weight_bits"), mapping=case["mapping_policy"],
                          real_registered_primitive=case.get("actual_registered_graph_primitive", False),
                          source_layer=case.get("source_id"), native=native_state, csim=csim_state,
                          hls=stages["hls"], route=stages["route"], bitstream=bit_state,
                          bitstream_sha256=bit_sha, static_interface_acceptance=interface_acceptance,
                          static_interface_evidence=interface_evidence,
                          board_acceptance=board_acceptance, board_execution=execution_state,
                          board_correctness="PASS_ALL_BATCH_FIXTURES" if fixture_ok and board_bound else "PASS_PROBE" if probe else "NO_ACCEPTED_RESULT",
                          board_timing=board_case.get("status", "NO_RESULT") if board_bound else "NOT_RUN" if probe else "NO_ACCEPTED_RESULT",
                          board_condition_deviations=board_summary.get("condition_deviations", []),
                          phase_clock_evidence_verified=board_summary.get("phase_clock_evidence_verified"),
                          hardware_measurement_eligible=board_summary.get("hardware_measurement_eligible"),
                          phase_evidence=phase_evidence,
                          hls_wall_seconds=hls_entry[1].get('wall_seconds') if hls_entry and stages['hls']=='PASS' else None,
                          route_wall_seconds=route_entry[1].get('wall_seconds') if route_entry and stages['route']=='PASS' else None,
                          board_runtime_seconds=runtime_seconds,probe_wall_seconds=probe_seconds,
                          board_runtime_source=board_runtime_path.as_posix() if runtime_rows else None,
                          probe_runtime_source=probe_runtime_path.as_posix() if probe_rows else None,
                          board_raw_csv_sha256=board_summary.get('raw_csv_sha256') if board_bound else None,
                          board_summary_sha256=sha(board_summary_path) if board_bound else None,
                          probe_correctness=probe,
                          board_dispatches=bound_dispatches,
                          board=board_status, fixture_count=len(case.get("fixtures", [])),
                          native_fixtures=record.get("result", {}).get("fixtures") if native_state == "PASS" else None,
                          native_mismatches=record.get("result", {}).get("mismatches") if native_state == "PASS" else None,
                          board_mismatches=sum(r.get("mismatched_elements", 0) for r in fixture_records) if fixture_records else 0 if probe else None,
                          case_sha256=case.get("case_sha256"),
                          hls_status_source=hls_entry[0].as_posix() if hls_entry else None,
                          route_status_source=route_entry[0].as_posix() if route_entry else None,
                          source=(folder / "case.json").as_posix(), formal_cost_eligible=False))
        for name, seconds in (("python_reference_prepare", case.get("preparation_wall_seconds")),
                              ("native_cpp_compile", record.get("compile_seconds")), ("native_cpp_run", record.get("run_seconds"))):
            stage_rows.append(dict(case_id=case_id, stage=name, seconds=seconds,
                                   status="RECORDED" if seconds is not None else "NOT_RUN", evidence_layer="SOFTWARE",
                                   source=(folder / "case.json" if name.startswith("python") else attempt / "native_summary.json").as_posix()))
        for name,entry in (('hls',hls_entry),('route',route_entry)):
            stage_rows.append(dict(case_id=case_id,stage=name+'_process_envelope_total',
                seconds=entry[1].get('wall_seconds') if entry and stages[name]!='STALE_SOURCE_BINDING' else None,
                status=stages[name],evidence_layer='PARENT_PROCESS_DO_NOT_ADD_CHILD_STAGES',
                source=entry[0].as_posix() if entry else None))
        paths = set()
        for base, kind in ((hls_dir, "HLS"), (route_dir, "ROUTE")):
            if base.exists():
                for name in ("hls_stages.jsonl", "shell_stages.jsonl", "stage_elapsed.jsonl"):
                    if kind == "ROUTE" and name == "hls_stages.jsonl":
                        continue  # Route exports repeat the original HLS log.
                    for path in base.rglob(name):
                        if path not in paths:
                            paths.add(path)
                            stage_rows.extend(_stages(path, case_id, kind, issues))
        xml = next((path for path in (hls_dir / "pilot_kernel_csynth.xml", hls_dir / "csynth.xml") if path.is_file()), None)
        hls_res = _hls_resources(xml) if xml else {}
        route_res = _route_resources(shell_path.parent / "utilization.rpt")
        route_core = _route_core_resources(shell_path.parent / "utilization_hier.rpt")
        resources.append(dict(case_id=case_id, hls_status=stages["hls"], route_status=stages["route"],
                              hls=hls_res, route=route_res, route_core=route_core, WNS_ns=shell.get("WNS_ns"),
                              WHS_ns=shell.get("WHS_ns"), WPWS_ns=shell.get("WPWS_ns"),
                              hls_source=xml.as_posix() if xml else None,
                              route_source=shell_path.as_posix() if shell else None))
        for scope in ("compute", "transfer", "call"):
            stats = board_case.get("scopes", {}).get(scope) or {}
            timing_rows.append(dict(case_id=case_id, scope=scope, status=board_status,
                                    warmup_n=board_case.get("warmup_rows"), measured_n=stats.get("n"),
                                    mean_ns=stats.get("mean_ns"), median_ns=stats.get("median_ns"),
                                    p95_ns=stats.get("p95_ns"), cv=stats.get("cv"),
                                    independent_batches=board_case.get("independent_batches"),
                                    measured_calls=board_case.get("measured_rows"),
                                    host_median_deviation=board_case.get("host_median_stability", {}).get("maximum_relative_deviation"),
                                    source=board_summary_path.as_posix() if board_summary else None))
        for batch in board_case.get("batches", []):
            governor_display = _cpu_governor_display(board_folder, batch["batch_id"])
            for scope in ("compute", "transfer", "call"):
                stats = batch.get("scopes", {}).get(scope) or {}
                batch_rows.append(dict(case_id=case_id, batch_id=batch["batch_id"], scope=scope,
                    status=batch.get("status"), metadata=batch.get("metadata", {}),
                    cpu_governor_display=governor_display["display"], cpu_governor_observations=governor_display,
                    measured_calls=batch.get("measured_rows"), warmup_n=batch.get("warmup_rows"), **stats))
        for label, source in [("ALL_BATCHES", board_case)] + [(b["batch_id"], b) for b in board_case.get("batches", [])]:
            for size, stats in source.get("host_scopes_by_group_size", {}).items():
                host_strata.append(dict(case_id=case_id, batch_id=label, group_size=int(size),
                    group_observations=source.get("host_group_size_distribution", {}).get(size), **(stats or {})))
    counts = inventory["counts_by_family"]
    extrap = extrapolate([], counts, physical_bindings_complete=False)
    conditional = _pilot_schedule(cases,inventory,attempt)
    extrap['status'] = 'NOT_ESTIMABLE'
    extrap['full_measurement_status'] = 'NOT_ESTIMABLE'
    extrap['conditional_pilot_schedule'] = conditional
    extrap['mapping_reduction_recommendation'] = conditional['mapping_recommendation']
    extrap["additional_missing_evidence"] = ["所有原语的实际物理映射绑定/去重",
        "P4之外各原语/规模/映射的代表性完整阶段耗时；网络传输与恢复等日历开销",
        "D09 CPU performance条件例外须在冻结前明确，或改用支持该模式的环境；当前偏差不倒改为合规",
        "当前整帧BRAM模板18键超过器件激活存储容量下界，4键尚无模板；实现覆盖须先闭合，不能仅缩减映射数量",
        "maxpool、stem普通3×3及更多空间尺度的代表项；D06若启用W16再追加独立计数",
        "每日可用板时、工具许可证和任务并行排程"]
    extrap["planning_scenarios"] = [dict(distinct_configurations_per_primitive=n, nominal_items=65*n,
        sequential_days_formula=f"{65*n} × t_mean_minutes / (60 × available_hours_per_day)",
        ideal_parallel_lower_bound_formula=f"{65*n} × t_mean_minutes / (60 × available_hours_per_day × independent_workers)",
        empirical_t_mean_minutes=None, includes_retries=True,
        note="公式情景，不是实测预测；t须由完整样本含失败重试求得，单板阶段仍串行，未覆盖类别必须补测") for n in (2, 3, 6)]
    annotations=_annotation_index(attempt,cases,issues)
    drivers=_driver_observations(attempt,cases,issues)
    return dict(inventory=inventory, cases=cases, resources=resources, timings=timing_rows, batches=batch_rows, host_strata=host_strata,
                mapping_comparisons=_mapping_comparisons(cases,timing_rows,resources), attempt=attempt.as_posix(),
                annotations=annotations,automatic_driver_observations=drivers,
                step_times=stage_rows, extrapolation=extrap, issues=issues, native_summary=native,
                instrumentation_audit=read_json(OUT / "instrumentation_audit.json", {}),
                clock_audit=read_json(OUT / "clock_configuration_audit.json", {}),
                board_environment_preflight=read_json(OUT / "preflight/board_transfer_readonly.json", {}),
                buffer_coverage=read_json(OUT / "buffer_coverage.json", {}),
                reset_interface_audit=reset_audit,
                original_workbook_audit=read_json(OUT / "original_workbook_update.json", {}),
                native_manifest_bound=bool(native_bound), preflight=read_json(OUT / "preflight/toolchain_access.json", {}))


def defaultdict_status(folder):
    found = {}
    paths = list(folder.rglob("status.json")) if folder.exists() else []
    # Fetches and board packages are immutable local copies. Do not scan board
    # probe result folders or promote a launch receipt to tool completion.
    if folder.name == "remote_evidence":
        packages = folder.parent / "board_packages"
        paths.extend(packages.glob("*/*/case/evidence/hls_status.json"))
        paths.extend(packages.glob("*/*/case/evidence/route_status.json"))
    for path in sorted(paths):
        record = read_json(path, {})
        key = (record.get("case_id"), record.get("stage"))
        if key[0] and key[1] in ("hls", "route"):
            previous = found.get(key)
            # Attempt directories are immutable and numbered; using both
            # start time and path supports RUNNING latest snapshots.
            rank = (record.get("started_utc", ""), str(path))
            if previous is None or rank > (previous[1].get("started_utc", ""), str(previous[0])):
                found[key] = (path, record)
    return found


def stage_audit(attempt, successor):
    """Write a local-only stage audit, without refreshing workbook/report.md.

    Completed tool process envelopes are counted exactly once. Their observed
    overlap proves the actual concurrency; extrapolation remains BUILD_ONLY,
    because route acceptance and board acceptance are separate evidence layers.
    """
    data = collect(attempt)
    latest = defaultdict_status(attempt / "remote_evidence")
    sources, envelopes, per_case = {}, [], {}
    expected = {row["case_id"]: row["case_sha256"] for row in data["cases"]}
    for (case_id, stage), (path, record) in latest.items():
        if record.get("case_sha256") != expected.get(case_id) or record.get("status") != "PASS":
            continue
        start = datetime.fromisoformat(record["started_utc"])
        finish = datetime.fromisoformat(record["finished_utc"])
        seconds = float(record["wall_seconds"])
        if finish <= start or seconds <= 0:
            raise ValueError("Invalid completed tool process envelope")
        envelopes.append(dict(case_id=case_id, stage=stage, started_utc=start.isoformat(),
                              finished_utc=finish.isoformat(), seconds=seconds))
        per_case[case_id] = per_case.get(case_id, 0) + seconds
        sources[path.as_posix()] = sha(path)
    events = sorted([(datetime.fromisoformat(row["started_utc"]), 1) for row in envelopes] +
                    [(datetime.fromisoformat(row["finished_utc"]), -1) for row in envelopes], key=lambda item: (item[0], item[1]))
    active, peak = 0, 0
    for _, delta in events:
        active += delta
        peak = max(peak, active)
    span = (events[-1][0] - events[0][0]).total_seconds() if events else None
    work_seconds = sum(row["seconds"] for row in envelopes)
    declared_jobs = read_json(attempt / "parallel_launch.json", {}).get("jobs")
    queue_cost_path = OUT / "preflight" / (attempt.name + "_observed_build_cost.json")
    queue_cost = read_json(queue_cost_path, {})
    queue_span = None
    if queue_cost:
        snapshot = Path(queue_cost.get("source_snapshot", ""))
        if not snapshot.is_file() or sha(snapshot) != queue_cost.get("source_snapshot_sha256"):
            raise ValueError("Observed queue cost is not bound to its source snapshot")
        queue_span = queue_cost.get("queue_wall_seconds")
        sources[queue_cost_path.as_posix()] = sha(queue_cost_path)
        sources[snapshot.as_posix()] = sha(snapshot)
    build_complete = len(envelopes) == 2 * len(data["cases"]) and all(row["route"] == row["hls"] == "PASS" for row in data["cases"])
    scenarios = []
    if build_complete and declared_jobs == peak == 3:
        for mappings in (2, 3, 6):
            count = data["inventory"]["unique_primitive_count"] * mappings
            predicted_seconds = span * count / len(data["cases"])
            schedule = repeated_build_mix(per_case, count, build_workers=3)
            scenarios.append(dict(distinct_mappings_assumed=mappings, nominal_physical_items=count,
                conditional_build_hours_at_observed_three_worker_throughput=predicted_seconds / 3600,
                conditional_build_days_at_24_hours_per_day=predicted_seconds / 86400,
                conditional_build_days_at_8_hours_per_day=predicted_seconds / (8 * 3600),
                conditional_three_worker_schedule={key: value for key, value in schedule.items() if key != "jobs"},
                full_measurement_calendar_days=None,
                assumptions="Same observed P4 build mix, three workers, similar machine/licenses and queue overhead; no board, repair, coverage redesign, or recovery time included"))
    successor_data = collect(successor)
    successor_state = dict(attempt=successor.name,
        statuses={name: dict(Counter(row[name] for row in successor_data["cases"])) for name in
            ("native", "csim", "hls", "route", "bitstream", "static_interface_acceptance", "board_execution", "board_correctness", "board_timing", "board_acceptance")},
        case_states=[{name: row[name] for name in ("case_id", "case_sha256", "hls", "route", "bitstream", "bitstream_sha256",
            "static_interface_acceptance", "static_interface_evidence", "hls_status_source", "route_status_source", "board_acceptance",
            "board_execution", "board_correctness", "board_timing", "probe_correctness")}
            for row in successor_data["cases"]],
        remote_toolchain_launch_receipt=read_json(successor / "remote_launch.json", {}),
        note="One fetched attempt05 package has independently bound route/bit/reset evidence. Static interface PASS is not board correctness or timing acceptance; launch receipts for remaining jobs are tool dispatch only.")
    for row in successor_data["cases"]:
        for value in (row["hls_status_source"], row["route_status_source"], row["static_interface_evidence"].get("hwh_source"),
                      row["probe_correctness"].get("summary_source"), row["probe_correctness"].get("snapshot_manifest_source")):
            if value:
                path = Path(value)
                sources[path.as_posix()] = sha(path)
    sources[(ROOT / "tools/hardware_measurement_pilot/board_runner.py").as_posix()] = sha(ROOT / "tools/hardware_measurement_pilot/board_runner.py")
    sources[(ROOT / "tools/hardware_measurement_pilot/schedule_scenarios.py").as_posix()] = sha(ROOT / "tools/hardware_measurement_pilot/schedule_scenarios.py")
    reset_path = OUT / "diagnostics/aux_reset_failure_20261012/reset_interface_audit.json"
    unreachable_path = OUT / "preflight/board_unreachable_diagnostic.json"
    for path in (reset_path, unreachable_path, attempt / "parallel_launch.json", successor / "source_manifest.json", successor / "native_summary.json"):
        if path.is_file():
            sources[path.as_posix()] = sha(path)
    result = dict(schema="pilot_stage_audit_v1", generated_utc=datetime.now(timezone.utc).isoformat(),
        scope="LOCAL_ARTIFACT_AUDIT_NO_NETWORK_NO_BOARD_NO_WORKBOOK_WRITE", attempt=attempt.name,
        statuses={name: dict(Counter(row[name] for row in data["cases"])) for name in
                  ("native", "csim", "hls", "route", "bitstream", "static_interface_acceptance", "board_execution", "board_correctness", "board_timing", "board_acceptance")},
        case_states=[{name: row[name] for name in ("case_id", "hls", "route", "bitstream", "bitstream_sha256",
                     "static_interface_acceptance", "board_execution", "board_correctness", "board_timing", "board_acceptance")} for row in data["cases"]],
        current_successor=successor_state,
        build_only_observations=dict(completed_process_envelopes=len(envelopes), completed_cases=len(per_case),
            declared_workers=declared_jobs, observed_peak_concurrent_envelopes=peak,
            first_started_utc=events[0][0].isoformat() if events else None,
            last_finished_utc=events[-1][0].isoformat() if events else None,
            observed_three_worker_queue_wall_seconds=queue_span,
            observed_three_worker_queue_started_utc=queue_cost.get("queue_started_utc"),
            observed_three_worker_queue_finished_utc=queue_cost.get("queue_finished_utc"),
            observed_calendar_span_seconds=span, summed_hls_route_process_seconds=work_seconds,
            ideal_work_divided_by_three_seconds=work_seconds / 3 if declared_jobs == 3 else None,
            worker_utilization_over_observed_span=work_seconds / (3 * span) if span and declared_jobs == 3 else None,
            case_process_seconds=per_case, process_envelopes=envelopes,
            duration_context="Concurrent shared-server process wall seconds, not isolated serial runtime or CPU-seconds. All-stage envelope includes PW-low HLS before the parallel queue; queue duration is separately reported.",
            double_counting_rule="Do not add child csim/HLS/synthesis/route/bitstream durations again to their parent process envelopes",
            acceptance="TOOL_BUILD_ONLY_RESET_INVALID_NO_ACCEPTED_BOARD_SAMPLE"),
        conditional_build_only_scenarios=scenarios,
        full_pipeline_estimate_status="NOT_ESTIMABLE_WITHOUT_ACCEPTED_BOARD_SAMPLES_AND_IMPLEMENTATION_COVERAGE",
        coverage_limit=dict(over_full_frame_bram_lower_bound=data["buffer_coverage"].get("exceeds_bram_lower_bound_count"),
                            no_template=data["buffer_coverage"].get("absent_template_count")),
        cost_accounting="attempt04 completed build work is retained as rejected-acceptance/rework cost; attempt05 and board recovery/testing are additional costs, never erase or reuse attempt04 acceptance",
        source_sha256=sources, formal_cost_eligible=False)
    target = OUT / ("stage_audit_" + attempt.name + "_reset.json")
    target.write_text(json.dumps(result, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")
    table = "\n".join(f"| {row['nominal_physical_items']} | {row['conditional_build_hours_at_observed_three_worker_throughput']:.3f} | {row['conditional_build_days_at_8_hours_per_day']:.3f} |" for row in scenarios)
    note = f"""# 实验3阶段审计：{attempt.name}的route事实与板端接纳分离

仅审阅本地归档，未联网、未访问板卡、未修改原工作簿或正式冻结选择，也未重新生成完整报告。

当前证据：{attempt.name}的14项本地整数/Vitis C仿真/HLS/route均PASS，bit生成及SHA单独保留；14项静态接口均为`RESET_INTERFACE_INVALID_NOT_BOARD_ACCEPTED`。低有效辅助复位接常数0的缺陷已由原始HWH确证，不能把route PASS升级成板端接纳。

首项body8_to_head16_reference_high的回执仅为DISPATCHED（PID 45778）。后续已归档诊断显示失联，未取回板端正确性或计时结果；不能据此宣称已执行到某个MMIO地址、已发生DMA或已永久损坏。其他13项无板端dispatch证据。

{successor.name}的本地整数状态为{successor_state['statuses']['native']}，新route状态为{successor_state['statuses']['route']}。当前已取回首项body8_to_head16_reference_high的真实包，HLS/route/bit哈希与完整静态route/reset门均通过（{successor_state['statuses']['static_interface_acceptance']}）。板端正确性状态为{successor_state['statuses']['board_correctness']}，计时状态为{successor_state['statuses']['board_timing']}。PASS_PROBE仅表示5组完整输出与整数参考一致，探测时的计数/单次耗时不导入latency列；至少3个独立批次的正式计时须由独立summary/CSV导入。其余项仅按已取回的本地证据登记，工具队列dispatch不提升为完成。

## 真实三并发的构建耗时

28个HLS/route进程包络的重叠峰值为{peak}，与队列jobs={declared_jobs}一致。三并发队列实际墙钟为{queue_span / 60 if queue_span else 0:.3f}分钟；包括队列启动前PW-low HLS在内，全部阶段首开始至末结束为{span / 60 if span else 0:.3f}分钟。重叠进程累计为{work_seconds / 60:.3f} job-minutes（进程墙钟和），不能称孤立顺序运行时间或CPU耗时；直接除3的理想下界为{work_seconds / 180:.3f}分钟。

新增排程函数将HLS与route作为三worker构建占用，再按构建完成顺序送入单板串行阶段。板端占用必须包含加载、各批校验和计时的完整独占时长，不能用单次推理latency代替，也不能除以3。由于当前没有已接纳板端样本，JSON中board占用、combined pipeline与full_measurement均为null；已记录的三worker构建情景仅供条件分析。排程器已用6项各10秒构建、各4秒板占用的可验算例验证：构建20秒、全部完成34秒。

以下仅是假定“新实现覆盖已补齐，任务分布与P4试做相近”的构建阶段条件情景，使用已观察的三并发整体吞吐；不含板端、恢复、改实现及重做成本，不是全量测量日历预测。

| 名义不同物理项数 | 仅构建小时 | 每天8小时的仅构建工作日 |
|---|---|---|
{table}

当前65键有18个超过整帧BRAM存储下界、4个缺模板，P4样本不能直接证明全量覆盖。attempt04虽完成构建但接纳失败，其实际耗时必须保留为失败/重做成本；attempt05及后续恢复、板测另外累加。只有通过新静态接口门槛并取得完整独立批次板测后，才有全流程外推样本。

完整状态、时间包络和来源SHA见[阶段审计JSON]({target.resolve().as_posix()})。
"""
    target.with_suffix(".md").write_text(note, encoding="utf-8")
    return result


def write_workbook(data, target):
    dependency_dir = OUT / "_deps"
    if dependency_dir.is_dir():
        sys.path.insert(0, str(dependency_dir))
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.comments import Comment
    wb = Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", fgColor="17365D")
    light_fill = PatternFill("solid", fgColor="EDF3F8")
    pending_fill = PatternFill("solid", fgColor="FFF2CC")

    def sheet(name, title, headers, rows, widths):
        ws = wb.create_sheet(name)
        ws.append([title])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        ws["A1"].font = Font(name="Microsoft YaHei", size=15, bold=True, color="17365D")
        ws.row_dimensions[1].height = 32
        ws.append(headers)
        for cell in ws[2]:
            cell.fill = header_fill
            cell.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF", size=10)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for row in rows:
            ws.append(list(row))
        for row in ws.iter_rows(min_row=3):
            for cell in row:
                cell.font = Font(name="Microsoft YaHei", size=10)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if cell.row % 2:
                    cell.fill = light_fill
                if isinstance(cell.value, float):
                    cell.number_format = "0.0000"
                if isinstance(cell.value, str) and cell.value.startswith(("NOT_", "PENDING", "DEFER", "STALE", "CANDIDATE")):
                    cell.fill = pending_fill
            ws.row_dimensions[row[0].row].height = 48 if name not in ("D09", "D07") else 96
        for index, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(index)].width = width
        ws.freeze_panes = "B3"
        ws.auto_filter.ref = f"A2:{get_column_letter(len(headers))}{ws.max_row}"
        ws.sheet_view.showGridLines = False
        ws.print_title_rows = "1:2"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.oddFooter.center.text = "实验3 · 试做结果不进入正式成本表 · 第 &P 页"
        return ws

    counts = Counter(row["native"] for row in data["cases"])
    count = len(data["cases"])
    fixture_count = sum(row["fixture_count"] for row in data["cases"])
    sheet("说明", "实验3 硬件测量流程试做：证据快照", ["项目", "值 / 说明", "来源"], [
        ("协议状态", "DRAFT_EXECUTE_PILOT_FREEZE_BEFORE_EXPERIMENT5", PROTOCOL_VERSION),
        ("计划", f"{count}项、每项5组唯一夹具，共{fixture_count}组；每项至少3个独立批次，每批50预热+至少1000正式调用", "inventory.json；case.json；D09"),
        ("当前本地C++", f"{counts.get('PASS',0)}/{count}项PASS；不代表Vitis C仿真或板上通过", "native_summary.json"),
        ("空值约定", "未测数值留空；状态明确NOT_RUN/PENDING，不把空值填0。", "report.py"),
        ("派生值", "来自可复现Python脚本的导入结果；此工作簿不含需要重算的公式。", "inventory.py；measurement.py；report.py"),
        ("原工作簿核查", "已按授权写入T21、W03/W04试做说明和D07/D09 M列；D06/D07/D09最终选择F、确认G、证据H仍为空，非正式冻结。审计状态：" + data["original_workbook_audit"].get("status", "NOT_CHECKED"), ORIGINAL_WORKBOOK.as_posix()),
        ("旧尝试", "attempt01的24项/120组本地验证留档，不再是当前硬件矩阵；W16、主体INT16、concat不进入当前批次。", "attempt01/reports_superseded_24case"),
    ], [24, 95, 65])
    sheet("D09", "D09 测量协议草案（执行本次pilot，实验5前冻结）", ["条目", "执行要求", "来源/证据"], D09, [25, 110, 70])
    sheet("D07", "D07 先2个不同物理点起步，按实测折中决定第3点", ["条目", "决策规则", "来源"], _d07_rows(data), [25, 110, 70])
    sheet("测试项", f"{count}项测试进度：软件与硬件证据分别记录",
          ["case_id", "真实已登记原语", "实际源层", "本地C++", "Vitis C仿真", "HLS综合", "布局布线", "bit", "板上", "夹具数", "本地不一致数", "板上不一致数"],
          [(r["case_id"], r["real_registered_primitive"], r["source_layer"], r["native"], r["csim"], r["hls"], r["route"], r["bitstream"], r["board"], r["fixture_count"], r["native_mismatches"], r["board_mismatches"]) for r in data["cases"]],
          [43, 18, 45, 20, 22, 20, 20, 28, 28, 10, 16, 16])
    sheet("板端接纳", "Route/bit生成事实与静态接口、板端正确性、计时接纳分别记录",
          ["case_id", "route事实", "bit生成与哈希", "bit SHA256", "静态接口接纳", "板端执行", "板端正确性", "板端计时", "板端接纳结论"],
          [(r["case_id"], r["route"], r["bitstream"], r["bitstream_sha256"], r["static_interface_acceptance"],
            r["board_execution"], r["board_correctness"], r["board_timing"], r["board_acceptance"]) for r in data["cases"]],
          [43, 18, 32, 70, 46, 60, 30, 30, 55])
    sheet("资源与时序", "核级HLS与包含shell的布局布线结果分列",
          ["case_id", "HLS LUT", "HLS FF", "HLS DSP", "HLS BRAM18", "集成route LUT", "集成route FF", "集成route DSP", "集成route BRAM18", "WNS ns", "WHS ns", "WPWS ns", "HLS状态", "route状态"],
          [(r["case_id"], *(r["hls"].get(k) for k in ("LUT", "FF", "DSP", "BRAM18")), *(r["route"].get(k) for k in ("LUT", "FF", "DSP", "BRAM18")), r["WNS_ns"], r["WHS_ns"], r["WPWS_ns"], r["hls_status"], r["route_status"]) for r in data["resources"]],
          [43, 14, 14, 14, 17, 18, 18, 18, 22, 14, 14, 14, 22, 22])
    sheet("route核资源", "从布线后层级报告的唯一core行读取；不使用总量相减估计",
          ["case_id", "core实例", "core LUT", "core FF", "core DSP", "core BRAM18等价", "route状态", "来源/空值说明"],
          [(r["case_id"], r["route_core"].get("instance"), *(r["route_core"].get(k) for k in ("LUT", "FF", "DSP", "BRAM18")),
            r["route_status"], "utilization_hier.rpt；缺失或表头/实例无法无歧义解析则留空") for r in data["resources"]],
          [43, 45, 18, 18, 18, 23, 25, 95])
    sheet('低高映射对照','同算子low/high实测；请求PE不等于实测吞吐，不能推断网络收益',
          ['算子','输入/权重位宽','配对状态','low compute ns','high compute ns','compute low/high','low transfer ns','high transfer ns',
           'low host中位数ns','high host中位数ns','low core LUT','high core LUT','low core FF','high core FF',
           'low core DSP','high core DSP','low core BRAM18','high core BRAM18'],
          [(r['operator'],f"I{r['input_bits']}/W{r['weight_bits'] or '-'}",r['status'],r['low']['compute_mean_ns'],r['high']['compute_mean_ns'],
            r['compute_low_over_high'],r['low']['transfer_mean_ns'],r['high']['transfer_mean_ns'],r['low']['host_median_ns'],r['high']['host_median_ns'],
            *(r[side]['core_resources'].get(resource) for resource in ('LUT','FF','DSP','BRAM18') for side in ('low','high')))
            for r in data['mapping_comparisons']], [28,20,24,22,22,20,22,22,25,25,18,18,18,18,18,18,22,22])
    ws = sheet("板上计时", "板上三范围计时；单位ns；预热不进入统计",
          ["case_id", "范围", "状态", "预热调用数", "统计观测数", "均值 ns", "中位数 ns", "P95 ns", "CV 样本SD/均值", "独立批次数", "正式调用数", "host批间中位数最大相对偏差", "原始结果来源"],
          [(r["case_id"], r["scope"], r["status"], r["warmup_n"], r["measured_n"], r["mean_ns"], r["median_ns"], r["p95_ns"], r["cv"], r["independent_batches"], r["measured_calls"], r["host_median_deviation"] if r["scope"] == "call" else None, r["source"]) for r in data["timings"]],
          [43, 15, 28, 14, 14, 20, 20, 20, 20, 14, 14, 28, 70])
    for cell in ws["I"][2:]:
        cell.number_format = "0.00%"
    sheet("逐批计时", "PL每调用一个观测；host每组均值只计一个观测；host不套5%CV门槛",
          ["case_id", "批次", "范围", "批次状态", "预热调用数", "正式调用数", "统计观测数", "均值 ns", "中位数 ns", "P95 ns", "CV", "实际PL起始Hz", "实际PL结束Hz", "CPU governor实际policy/原标记", "起始温度C", "结束温度C", "独立事件", "起始UTC", "结束UTC"],
          [(r["case_id"], r["batch_id"], r["scope"], r["status"], r["warmup_n"], r["measured_calls"], r.get("n"), r.get("mean_ns"), r.get("median_ns"), r.get("p95_ns"), r.get("cv"),
            *(r["metadata"].get(k) for k in ("clock_hz", "clock_hz_end")), r.get("cpu_governor_display"),
            *(r["metadata"].get(k) for k in ("temperature_c", "temperature_c_end", "independence_event", "started_utc", "finished_utc"))) for r in data["batches"]],
          [43, 18, 15, 20, 14, 14, 14, 20, 20, 20, 15, 20, 20, 20, 16, 16, 22, 30, 30])
    sheet("host按N分层", "不同N的组均值分别统计；每组一个观测，不能视为N个独立样本",
          ["case_id", "批次或汇总", "N", "组观测数", "均值 ns", "中位数 ns", "P95 ns", "CV（仅描述）"],
          [(r["case_id"], r["batch_id"], r["group_size"], r["group_observations"], r.get("mean_ns"), r.get("median_ns"), r.get("p95_ns"), r.get("cv")) for r in data["host_strata"]],
          [43, 25, 12, 16, 22, 22, 22, 22])
    audit = data["instrumentation_audit"]
    sheet("计时仪表审计", "历史失败独立记录；数值正确不能替代相位语义正确", ["项目", "状态/发现", "证据"], [
          ("历史attempt01相位审计", audit.get("status", "NOT_CHECKED"), str(audit.get("evidence", {}))),
          ("历史attempt01整数C仿真", audit.get("numeric_csim_status", "NOT_RUN"), "仅数值一致性"),
          ("历史attempt01问题", "；".join(audit.get("reason", [])), "instrumentation_audit.json"),
          ("当前相位证据", f"所选attempt有{sum(r['phase_evidence']['status']=='PASS_BOUND_PHASE_EVIDENCE' for r in data['cases'])}项绑定生成FSM/AXI adapter静态检查、完整时钟/复位、timer RTL与板端计数；完整动态集成RTL=NOT_RUN", "每case phase_validation.json与phase_evidence目录；只接纳片上计算区间含局部控制开销"),
          ("attempt03时钟配置", "请求200MHz，实际HWH为187.5MHz，偏差-6.25%；旧队列停止并保留耗时，attempt04修订一致", "clock_configuration_audit.json"),
          ("attempt05时钟", "PL配置187.5MHz、route周期16/3ns，HLS为更严格的5ns；板端读取寄存器派生实际频率", "attempt05 shell_metadata.json；Clocks.fclk0_mhz；逐批环境记录"),
          ("独立物理频率校准", "NOT_RUN；当前不声称独立校准，寄存器派生频率和主机区间sanity分别记录", "D09草稿"),
          ("本次CPU条件", "逐批policy实际读值userspace，仅支持userspace；原summary UNAVAILABLE表示performance控制不可用。保留protocol_conditions_met=false；冻结前明确例外或更换环境，当前不倒改PASS", "选定run/batch*/environment_start.json、environment_end.json及原summary"),
          ("attempt04辅助复位缺陷", "14个case的原始包辅助reset低有效却接0；RESET_INTERFACE_INVALID_NOT_BOARD_ACCEPTED。route/timing PASS事实保留。", "diagnostics/aux_reset_failure_20261012/reset_interface_audit.json"),
          ("首次probe证据边界", "只有DISPATCHED回执；后续归档诊断失联，无板端正确性/计时结果；不能从PID推断执行到具体MMIO/DMA阶段。", "attempt04/board_transport/body8_high_probe01/receipt_launch.json；preflight/board_unreachable_diagnostic.json"),
    ], [32, 105, 100])
    sheet("阶段耗时", "父级进程总时与子阶段分列，不能重复相加；board为进程占用而非单次latency",
          ["case_id", "阶段", "墙钟秒", "状态", "证据层", "来源"],
          [(r["case_id"], r["stage"], r["seconds"], r["status"], r["evidence_layer"], r["source"]) for r in data["step_times"]], [43, 34, 18, 20, 20, 100])
    inv = data["inventory"]
    sheet("全量外推", "完整外推尚未成立；名义配置数不等同于实测物理项数", ["项目", "值", "证据/限制"], [
        ("本次W8A8/W8A16域去重原语", inv["unique_primitive_count"], "72架构×当前2精度域枚举"),
        ("含D06待确认W16的全登记域", inv["conditional_all_registered_profiles_unique_primitive_count"], "仅条件候选数，不表示启用W16"),
        ("实际不同物理实现数", None, "待绑定PE/有效SIMD/tile/buffer/接口/时钟/源码并去重"),
        ("每原语2种独立配置的名义项数", inv["nominal_full_measurement_items_by_physical_configurations_per_primitive"]["2"], "条件示例；不是已冻结物理测量总数"),
        ("每原语3种独立配置的名义项数", inv["nominal_full_measurement_items_by_physical_configurations_per_primitive"]["3"], "条件示例；不是已冻结物理测量总数"),
        ("每原语6种独立配置的名义项数", inv["nominal_full_measurement_items_by_physical_configurations_per_primitive"]["6"], "六个策略名可能选择同一物理行"),
        ("完整全量可信日历", None, "NOT_ESTIMABLE；覆盖、物理去重与日历开销未闭合"),
        ("完整全量是否超过21天", None, "不把P4条件排程等同于全65键承诺"),
        ("已测工具+板端worker阶段是否超过21个8h条件工作日", data['extrapolation']['conditional_pilot_schedule']['conditional_time_trigger_over_21_workdays'], "仅当前实测mix和已记录worker占用；完整流程含桥接/人工/修复的日期仍为空"),
        ("当前建议", "2个不同物理点起步，必要时第3非支配点", "六策略可以复用物理行；详见D07条件决策"),
        ("当前整帧模板存储下界超容量键数", data["buffer_coverage"].get("exceeds_bram_lower_bound_count"), "分析下界，不是route结果；18个键需要不同实现后独立测量"),
        ("当前无模板的结构键数", data["buffer_coverage"].get("absent_template_count"), "普通3×3和maxpool共4个；剩余键也未因此证明fit"),
        ("已观测完整P4过程项数", data['extrapolation']['conditional_pilot_schedule']['observed_case_count'], "需要14项完整mix才生成数值排程；CPU偏差不会抹去实际worker占用"),
        *[("缺少证据", item, "需实测/绑定") for item in data["extrapolation"]["additional_missing_evidence"]],
    ], [44, 68, 85])
    sheet('P4条件排程','65×2/3/6仅为P4 mix重复条件；3构建worker＋单板，保留userspace偏差',
          ['假设不同配置数','名义项数','仅构建h','单板累计占用h','构建至单板完成h','已测工具+板端worker阶段条件工作日（8h/日）','已测工具+板端worker阶段条件工作日（24h/日）','完整全量日期','未计开销'],
          [(r['distinct_physical_configurations_assumed'],r['nominal_items'],r['conditional_build_hours'],r['single_board_runtime_hours'],
            r['conditional_pipeline_hours'],r['conditional_workdays_at_8h'],r['conditional_workdays_at_24h'],None,
            '网络/启动、人工复核、恢复、额外失败重试与改实现；覆盖18 BRAM/4模板及物理去重未闭合')
            for r in data['extrapolation']['conditional_pilot_schedule']['scenarios']], [20,18,20,22,25,48,48,24,100])
    sheet('自动driver实耗','自动案例从probe上传开始至timing导入结束；排除首个手工案例与未完成项',
          ['case_id','driver总秒数','probe+timing worker秒数','worker外合计残差秒数','来源','范围限制'],
          [(r['case_id'],r['driver_elapsed_seconds'],r['probe_plus_timing_worker_seconds'],r['outside_worker_residual_seconds'],r['source'],
            '残差含桥接/轮询/本地复核/编排，非纯网络耗时；未外推到14项或全量排程') for r in data['automatic_driver_observations']['records']],
          [43,24,30,30,100,100])
    sheet('异常派生标注','全部原始行保留，剔除数0；MAD0标签不等于稳健异常',
          ['case_id','输入行数','输出行数','invalid行','带任意标签行','host高时延标签行','PL高时延标签数（scope相加）','MAD0描述标签数（scope相加）','注释CSV','规则JSON'],
          [(r['case_id'],r['input_rows'],r['output_rows'],r['invalid_rows'],r['rows_with_any_flag'],
            r['flag_row_counts'].get('CALL_DESCRIPTIVE_HIGH_LATENCY',0),
            sum(r['flag_row_counts'].get(scope+'_DESCRIPTIVE_HIGH_LATENCY',0) for scope in ('COMPUTE','TRANSFER')),
            sum(value for key,value in r['flag_row_counts'].items() if 'MAD_ZERO_NO_ROBUST_DECISION' in key),r['annotated_csv'],r['rules'])
           for r in data['annotations']['records']], [43,16,16,16,22,24,24,34,100,100])
    sheet("原语盘点", "全量去重类别及精度：当前网络真实结构",
          ["类别", "输入位宽", "权重位宽", "输出位宽", "结构/精度种数"],
          [(r["family"], r["input_bits"], r["weight_bits"], r["output_bits"], r["count"]) for r in inv["counts_by_family_precision"]], [32, 20, 20, 20, 24])
    sheet("来源绑定", "原始资料和生成器SHA256", ["来源", "SHA256", "说明"],
          [(name, value, "库存枚举使用的源码") for name, value in inv["source_sha256"].items()] +
          [("tools/hardware_measurement_pilot/report.py", sha(__file__), "本工作簿生成器")], [92, 72, 40])
    wb.save(target)
    loaded = load_workbook(target, data_only=False)
    errors, formulas = [], []
    for ws in loaded:
        for row in ws:
            for cell in row:
                if cell.data_type == "e":
                    errors.append(f"{ws.title}!{cell.coordinate}: {cell.value}")
                if cell.data_type == "f":
                    formulas.append(f"{ws.title}!{cell.coordinate}")
    if errors or formulas:
        raise RuntimeError("Unexpected spreadsheet errors/formulas: " + repr(errors + formulas))
    assert loaded["测试项"].max_row == count + 2 and loaded["板上计时"].max_row == count * 3 + 2
    return dict(status="PASS", sheets=loaded.sheetnames, case_rows=count, timing_rows=count*3,
                formula_cells=0, excel_errors=[], unavailable_numeric_cells="BLANK", workbook_sha256=sha(target))


def update_original_workbook():
    """Apply only the user-authorized T21/W03/W04 and D07/D09 note edits.

    Keep final-decision F/G/H cells untouched, preserve all existing formulas,
    styles and unrelated values, retain a byte-identical original backup, and
    compare the original SHA immediately before replacing the reviewed file.
    """
    sys.path.insert(0, str(OUT / "_deps"))
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment
    audit_path = OUT / "original_workbook_update.json"
    previous = read_json(audit_path, {})
    current_sha = sha(ORIGINAL_WORKBOOK)
    if previous.get("after_sha256") == current_sha:
        return dict(previous, action="ALREADY_APPLIED_NO_CHANGE")
    expected_sha = "95323753df58cfc09d295659a64cd0abe47d48c6df294805fd99178e0604f4fb"
    if current_sha != expected_sha:
        raise ValueError("Original workbook changed since exact-cell review; re-read before applying edits")
    wb = load_workbook(ORIGINAL_WORKBOOK)
    all_t21_hits = [(s.title, c.coordinate, str(c.value)) for s in wb for row in s for c in row
                    if c.value is not None and "T21" in str(c.value)]
    if all_t21_hits:
        raise ValueError("T21 exists; locate and review instead of appending: " + repr(all_t21_hits))
    original_cells = {(s.title, c.coordinate): (c.value, copy(c._style), c.hyperlink.target if c.hyperlink else None,
                                               c.hyperlink.location if c.hyperlink else None)
                      for s in wb for row in s for c in row}
    original_sheets = list(wb.sheetnames)
    merges = {s.title: list(map(str, s.merged_cells.ranges)) for s in wb}
    assert wb["先确定事项"]["A12"].value == "D07"
    assert wb["先确定事项"]["A14"].value == "D09"
    assert wb["执行顺序与并行"]["A9"].value == "W03"
    assert wb["执行顺序与并行"]["A10"].value == "W04"
    mapping = ("14项试做：PW INT8及头部I16/W8/O16各low PE1/SIMD2、high PE3/SIMD4（共4项）；"
               "DW3、DW5各low PE2/SIMD1、high PE4/SIMD1（共4项）；"
               "nearest2x INT8、requant add INT8、body8→head16边界各low PE2/SIMD1、high PE4/SIMD1（共6项）。"
               "无权重操作以PE通道并行，SIMD固定1。W16权重待D06确认；concat不列入默认试做。")
    draft = ("DRAFT，按草案执行本次pilot，实验5前正式冻结；不填写F/G/H正式冻结栏。"
             "至少3个独立批次，每批重新加载bit或重启runtime、50次预热、≥1000次正式调用。"
             "输出拷贝计入call计时，逐元素比较/hash在计时外；短算子分组N次计总时/N，保留实际N与每次PL计数。"
             "PL compute/transfer样本CV≤5%；host报告批间中位数稳定性与P95，不预设未经确认的硬阈值。"
             "逐批记录实际PL时钟、CPU governor和起止温度；有可用接口且支持时设performance并恢复原值，缺失如实标记限制。"
             "异常/超时记录保留并标记，不删除后重跑至通过。")
    allowed = set()

    def put(sheet, coordinate, value):
        allowed.add((sheet, coordinate))
        wb[sheet][coordinate] = value

    decision = wb["先确定事项"]
    for coordinate, reference in (("M5", "L5"), ("M12", "L12"), ("M14", "L14")):
        decision[coordinate]._style = copy(decision[reference]._style)
        decision[coordinate].alignment = Alignment(wrap_text=True, vertical="top")
    put("先确定事项", "M5", "试做补充说明（非冻结）")
    put("先确定事项", "M12", "试做范围，非正式冻结")
    put("先确定事项", "M14", draft + "\n协议：" + (OUT / "D09_protocol.md").as_posix())
    decision["M14"].hyperlink = (OUT / "D09_protocol.md").as_posix()
    decision.column_dimensions["M"].width = 100
    put("执行顺序与并行", "E9", str(wb["执行顺序与并行"]["E9"].value) + "\n" + mapping + "仅试做，不据此冻结正式总项数；详见T21。")
    put("执行顺序与并行", "D10", "每键W03通过；本次pilot按D09 DRAFT执行，实验5前完成正式冻结；正式测量仍须冻结协议。")
    put("执行顺序与并行", "E10", str(wb["执行顺序与并行"]["E10"].value) + "\n" + mapping + "\n" + draft)
    task_sheet = wb["待办优先级"]
    row_number = task_sheet.max_row + 1
    values = ["T21", "P0", "实验3 硬件测量流程试做（14项）", "S2-03至S2-12；D06/D07/D09",
              mapping + "\n" + draft + "\n逐步记录实际耗时；正式物理项数和映射缩减等待完整外推，不把六策略名直接乘为六种实现。",
              "按当前用户指定试做范围；逐项整数合同/HLS门槛通过后流转",
              "逐项整数/CSIM/HLS/route/bit/board证据、三批原始计时、阶段耗时、D09草案与D07条件建议",
              "仅流程pilot，不授予正式成本表/NAS准入", "FPGA/测量", "试做执行与计时仪表修订；未验收",
              None, None, "用户2026-10-12确认试做范围及D09修订；results/hardware_measurement_pilot_20261012",
              "W03/W04", "各项可独立流水；同板计时不并行干扰；原始异常保留", "D09 DRAFT；实验5前正式冻结，D07试做非正式冻结"]
    for column, value in enumerate(values, 1):
        cell = task_sheet.cell(row_number, column)
        cell._style = copy(task_sheet.cell(row_number-1, column)._style)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        put(task_sheet.title, cell.coordinate, value)
    task_sheet.row_dimensions[row_number].height = 300
    wb["执行顺序与并行"].row_dimensions[9].height = max(wb["执行顺序与并行"].row_dimensions[9].height or 15, 160)
    wb["执行顺序与并行"].row_dimensions[10].height = max(wb["执行顺序与并行"].row_dimensions[10].height or 15, 300)
    backup_dir = OUT / "workbook_backup"
    backup_dir.mkdir(parents=True, exist_ok=True)
    backup = backup_dir / (ORIGINAL_WORKBOOK.stem + "_before_20261012_" + expected_sha[:12] + ".xlsx")
    if backup.exists():
        if sha(backup) != expected_sha:
            raise ValueError("existing backup is not the reviewed workbook")
    else:
        shutil.copy2(ORIGINAL_WORKBOOK, backup)
    temporary = OUT / "original_workbook_updated_candidate.xlsx"
    wb.save(temporary)
    checked = load_workbook(temporary)
    assert checked.sheetnames == original_sheets
    assert {s.title: list(map(str, s.merged_cells.ranges)) for s in checked} == merges
    unexpected = []
    for (name, coordinate), before in original_cells.items():
        if (name, coordinate) in allowed:
            continue
        cell = checked[name][coordinate]
        after = (cell.value, cell._style, cell.hyperlink.target if cell.hyperlink else None,
                 cell.hyperlink.location if cell.hyperlink else None)
        if before != after:
            unexpected.append(name + "!" + coordinate)
    if unexpected:
        raise ValueError("Unrelated workbook cells changed: " + repr(unexpected))
    for row in (11, 12, 14):
        assert all(checked["先确定事项"].cell(row, column).value is None for column in (6, 7, 8))
    if sha(ORIGINAL_WORKBOOK) != expected_sha:
        raise ValueError("original changed during editing; not replacing it")
    # A fully verified copy replaces only the explicitly authorized workbook.
    shutil.copy2(temporary, ORIGINAL_WORKBOOK)
    after_sha = sha(ORIGINAL_WORKBOOK)
    changes = [{"sheet": name, "cell": coordinate,
                "before": original_cells.get((name, coordinate), (None,))[0],
                "after": checked[name][coordinate].value} for name, coordinate in sorted(allowed)]
    result = dict(status="PASS_AUTHORIZED_SCOPED_UPDATE", date="2026-10-12", original=ORIGINAL_WORKBOOK.as_posix(),
                  before_sha256=expected_sha, after_sha256=after_sha, backup=backup.as_posix(),
                  t21_preexisting_hits=all_t21_hits, t21_row=row_number, changes=changes,
                  unrelated_cells_unchanged=True, existing_formulas_styles_hyperlinks_unchanged=True,
                  D06_D07_D09_final_F_G_H_unchanged=True, formal_freeze_written=False)
    audit_path.write_text(json.dumps(result, ensure_ascii=False, indent=2, allow_nan=False)+"\n", encoding="utf-8")
    return result


def generate(attempt):
    attempt = attempt.resolve()
    if not attempt.is_relative_to(OUT.resolve()) or not attempt.is_dir():
        raise ValueError("attempt must be an existing child of the pilot results directory")
    data = collect(attempt)
    timestamp = datetime.now(timezone(timedelta(hours=8))).isoformat(timespec="seconds")
    cases = data["cases"]
    case_count = len(cases)
    planned_fixture_count = sum(row["fixture_count"] for row in cases)
    native_pass = sum(r["native"] == "PASS" for r in cases)
    fixtures = sum(r["native_fixtures"] or 0 for r in cases)
    table = "\n".join("| " + " | ".join(str(row[key]) for key in ("case_id", "native", "csim", "hls", "route", "bitstream", "board")) + " |" for row in cases)
    acceptance_table = "\n".join("| " + " | ".join(str(row[key]) for key in
        ("case_id", "static_interface_acceptance", "board_execution", "board_correctness", "board_timing", "board_acceptance")) + " |" for row in cases)
    invalid_interfaces = sum(row["static_interface_acceptance"] == "RESET_INTERFACE_INVALID_NOT_BOARD_ACCEPTED" for row in cases)
    runtime_pass=sum(row['board_correctness']=='PASS_ALL_BATCH_FIXTURES' and row['board_timing']=='PASS' for row in cases)
    cpu_deviation=sum(row['board']=='STATISTICS_PASS_WITH_CPU_GOVERNOR_DEVIATION' for row in cases)
    compute_records=[row for row in data['timings'] if row['scope']=='compute']
    measured_calls=sum(row['measured_calls'] or 0 for row in compute_records)
    warmup_calls=sum(row['warmup_n'] or 0 for row in compute_records)
    batches=sum(row['independent_batches'] or 0 for row in compute_records)
    board_note=(f"当前绑定板端结果：{runtime_pass}/{case_count}项三批夹具与原始数值统计通过，累计{batches}批、"
        f"{measured_calls}次正式调用和{warmup_calls}次预热；其中{cpu_deviation}项保留CPU governor条件偏差。"
        '运行事实、协议条件及正式成本资格分别报告；历史连接预检不作为板上正确性依据。')
    plan=data['extrapolation']['conditional_pilot_schedule']
    schedule_table=_schedule_markdown(data)
    comparison_table=_comparison_markdown(data)
    boundary=next((row for row in data['mapping_comparisons'] if row['operator']=='body8_to_head16' and row['status']=='OBSERVED_PAIR'),None)
    negative_note=(f"body8→head16的compute均值：low为{boundary['low']['compute_mean_ns']/1000:.6f}μs，"
        f"high为{boundary['high']['compute_mean_ns']/1000:.6f}μs，low/high={boundary['compute_low_over_high']:.6f}。"
        '请求PE增加未带来这一边界算子的实测计算加速，应保留该负结果；不能把PE翻倍当成吞吐翻倍。'
        if boundary else '元素算子的high/low是否获益须等待同算子配对实测；请求PE不能直接替代吞吐数据。')
    annotation_links='\n'.join(f"| {row['case_id']} | {row['input_rows']} | {row['output_rows']} | {row['invalid_rows']} | {link(row['annotated_csv'],'注释CSV')} | {link(row['rules'],'规则与哈希')} |"
        for row in data['annotations']['records'])
    report = f"""# 实验3 硬件测量流程试做报告

生成时间：{timestamp}。本次尝试：{link(attempt)}。

**当前矩阵为{case_count}项W8A8/W8A16及混合精度边界；本地C++整数一致性验证 {native_pass}/{case_count} 项、{fixtures}/{planned_fixture_count} 组夹具。已绑定板端三批夹具与原始数值统计通过{runtime_pass}/{case_count}项；CPU条件偏差单列，试做结果不进入正式成本表。**

当前已绑定产物中，{invalid_interfaces}项为`RESET_INTERFACE_INVALID_NOT_BOARD_ACCEPTED`。route/timing PASS和bit已生成仅保留为工具事实；静态接口与板端接纳分别判断，后台DISPATCHED回执不能充当正确性或计时结果。

已按用户授权更新{link(ORIGINAL_WORKBOOK)}：新增T21（待办优先级第23行），在W03/W04写明14项及实际PE/SIMD；D07的M列写“试做范围，非正式冻结”，D09的M列写草案说明及协议链接。“先确定事项”D06/D07/D09第11/12/14行的F（最终选择）、G（确认人/日期）、H（证据）仍为空。本次按D09草案执行，实验5前正式冻结。原文件备份及语义复核见{link(OUT/'original_workbook_update.json')}，状态为{data['original_workbook_audit'].get('status','NOT_CHECKED')}；WPS保存的空值与格式序列化规范化另有记录。

{board_note}

## 样本与数值范围

实际主线为elastic_neck72_v1，使用真实P4形状：PW头部48×40×40→21×40×40；DW3/DW5为144×40×40；nearest为48×20×20→48×40×40；add为两路48×40×40。上述INT8操作共10个映射项；W8A16头部PW另有2项；48×40×40的body8→head16重定标边界另有2项。DW/nearest/add只测INT8，权重W16推迟到D06确认；当前网络与工作簿没有确认concat必要性，因此默认不纳入。每项5组夹具覆盖随机值、激活、ties-even、溢出与饱和。该流程使用合成整数输入/权重，不声称训练权重部署或检测精度。

旧attempt01的24项/120组本地软件探针保留在{link(OUT/'attempt01/reports_superseded_24case')}；其W16、主体INT16和concat不能升级为本次已授权硬件矩阵，也不能混入本次进度分母。

源码依据：{link(ROOT/'tools/modnas_elastic/cost_graph.py')}、{link(ROOT/'tools/fpga_h4/quant_contracts.py')}、{link(attempt/'inventory.json')}、{link(attempt/'source_manifest.json')}。

## 每项证据状态

| 测试项 | 本地C++ | Vitis C仿真 | HLS | 布局布线 | bit | 板上 |
|---|---|---|---|---|---|---|
{table}

| 测试项 | 静态接口接纳 | 板端执行 | 板端正确性 | 板端计时 | 板端接纳结论 |
|---|---|---|---|---|---|
{acceptance_table}

逐项LUT（查找表）、FF（触发器）、DSP（数字信号处理单元）、BRAM18（18Kb块RAM）、WNS（最差建立时间裕量）及三个板上时间范围见工作簿。HLS核级资源与含AXI/计数器shell的布线后资源分别保存。本地C++的运行秒数只用于流程耗时，不能填为FPGA延迟。

## 同算子high/low实测对照

{comparison_table}

{negative_note} DW等算子的配对比值与core LUT/FF/DSP/BRAM一并列在工作簿“低高映射对照”；host中位数包含软件调用开销，不能与PL compute混为一项。这些是相同P4算子、当前输入与实现的观测，不外推为全网络加速或硬件收益。

## 计时仪表与D09执行边界

历史attempt01的整数C仿真与HLS综合通过，但生成RTL把相位2/3安排到计算结束后相邻状态，且空闲相位寄存器保留3，计时仪表判为FAIL_INSTRUMENTATION_SEMANTICS。证据见{link(OUT/'instrumentation_audit.json')}。该负结果保留；当前{sum(row['phase_evidence']['status']=='PASS_BOUND_PHASE_EVIDENCE' for row in cases)}项另有逐case、源码、bit与HWH绑定的phase_validation证据，不继承旧数值PASS。

当前phase2仅称**片上计算区间（含局部控制开销）**。证据链是生成FSM/AXI adapter静态检查、完整HWH时钟/复位连接、独立timer RTL测试和五夹具板端计数；完整动态集成RTL仿真明确为NOT_RUN。不把该区间扩写为仅算术指令周期，也不把phase1+3称为DDR总线占用率。逐项路径保存在{link(OUT/'report_data.json')}的phase_evidence。

attempt03请求200MHz，但Vivado实际分频与HWH均为187.5MHz（-6.25%），因此停止旧队列并保留失败准备成本。attempt04统一时钟后又暴露辅助复位低有效接0的缺陷，14项route通过但未被板端接纳；attempt05修正完整复位连接，PL配置为187.5MHz、route周期16/3ns，HLS的5ns仍单列为更严格综合目标。详见{link(OUT/'clock_configuration_audit.json')}和{link(OUT/'stage_audit_attempt04_reset.json')}。板上周期转换使用PYNQ Clocks从寄存器派生的实际频率，每批及每组前后检查稳定性，并检查PL总区间包含于主机调用区间；独立物理频率校准为NOT_RUN，不作已校准声明。

逐批真实环境读到policy0为userspace、可用模式仅userspace；scaling_cur_freq原始记录为1199999kHz。冻结runner把performance控制不可用粗标为UNAVAILABLE，报告同时保留这个原值和已读到的实际policy，不能解释为governor无法读取。实际温度来自逐批IIO/hwmon记录。当前CPU偏差使protocol_conditions_met=false、hardware_measurement_eligible=false，但不抹去夹具、计时与相位验证通过事实。实验5前应明确冻结performance例外或更换支持该模式的环境；当前不倒改原summary为合规PASS。

每项至少3个独立批次，每批重载bit或重启runtime，50次预热后至少1000次正式调用。输出拷贝包含在call时间，逐元素比较与hash在计时窗外；短算子可按N次分组，总时/N，每次PL计数仍保留，host每组均值只计一次统计观测。PL compute/transfer逐批及合并CV≤5%；host报告批间中位数稳定性和P95，不默认套用5% CV或未经确认的中位数硬阈值。每批保存实际PL起止频率、CPU governor和起止温度；不可用接口如实记条件偏差。

## 异常保留与派生标注

{link(data['annotations']['source'],'选定run异常标注索引')}当前覆盖{data['annotations']['case_count']}项、{data['annotations']['total_rows']}个原始行。原始行全部保留状态为{data['annotations']['all_raw_rows_preserved']}，统计剔除数为{data['annotations']['statistical_exclusions']}。规则按case/batch/预热或正式/scope分层，以median+6×1.4826×MAD作描述性高时延标签；host先按唯一group形成一个总时/N观测，再映射回成员行。MAD=0时明确不作稳健异常判定，“高于中位数”仅是另一种描述标签。invalid、慢样本和原顺序均保留；注释不修改统计、验收阈值或选样。

| 测试项 | 输入行数 | 输出行数 | invalid行 | 注释 | 规则 |
|---|---|---|---|---|---|
{annotation_links}

## 阶段耗时与全量外推

原始本地准备记录：{link(attempt/'preparation.json')}；本地C++记录：{link(attempt/'native_summary.json')}。已测阶段按项导入“阶段耗时”页，未执行阶段为空。每个HLS/route尝试的日志及阶段秒数分别保存，不删除失败项。

当前W8A8/HEAD_W8A16结构/精度盘点为65种原语：{json.dumps(data['inventory']['counts_by_family'],ensure_ascii=False)}。65×2/3/6得到130/195/390只在每种原语确有对应数量的**不同物理配置**时成立；六个映射策略名可重用同一行，实际物理项数尚未绑定。若D06后来确认启用HEAD_W16A16，全登记域才扩展为69种；该条件候选数不是当前启用范围。

**完整65键全量可信日期仍为NOT_ESTIMABLE；已观测P4过程支持下列有限条件排程。** 将当前14项low/high mix依固定次序重复到130/195/390个名义物理项，使用各项实际HLS+route进程秒数、3个构建worker及单板串行probe＋三批timing进程占用。CPU userspace条件偏差保留，实际占用不因未获D09条件资格而消失。构建任务完成后才进入单板队列，板端时间不除以3。

{schedule_table}

其中单板占用含worker内加载、夹具检查和全部预热/计时，不等于单次推理latency；网络上传/下载及worker外启动、人工复核、空闲排队、恢复、额外失败重试和改实现时间未包含。构建进程耗时来自当时共享机器上下文，不能称孤立顺序基准。8h/24h数值仅称“已测工具+板端worker阶段条件工作日”；完整流程（含桥接、人工复核、实现修复）的日期仍为空。attempt04的已耗构建仍作为失败接纳/重做成本保留，不能抹去或当成05的合格结果。原始依据、假设及缺项见{link(OUT/'extrapolation.json')}。

另有{data['automatic_driver_observations']['case_count']}个已完成自动案例的真实driver包络，从probe上传开始至timing导入结束：均值{_number(data['automatic_driver_observations']['driver_elapsed']['mean_seconds'],3)}秒，范围{_number(data['automatic_driver_observations']['driver_elapsed']['minimum_seconds'],3)}–{_number(data['automatic_driver_observations']['driver_elapsed']['maximum_seconds'],3)}秒。减去probe＋timing worker后的残差均值为{_number(data['automatic_driver_observations']['outside_worker_residual']['mean_seconds'],3)}秒，含桥接、轮询、本地校验及编排，不能称纯网络耗时。这些记录不混入首个手工body8-high案例，也不为凑齐14项补数；本段只描述真实开销，未加入上述条件排程。逐项事件与哈希见工作簿“自动driver实耗”及report_data.json。

{_d07_rows(data)[1][1]} 因此建议先2种不同物理配置起步，若需要未覆盖的实测非支配折中再增加第3种；六个查询策略可以复用这些物理行。

另一个实质限制是当前内核将完整输入/输出帧显式绑定到BRAM：65个当前结构键中有18个仅激活存储下界就超过624 BRAM18，另4个缺少算子模板；剩余43个也只是未触发该下界，尚非资源可行证明。例如INT8 1×1卷积16×320×320→48×320×320仅两帧即6.25MiB，至少需要2845 BRAM18，而器件容量约1.371MiB。分析及独立复核见{link(OUT/'buffer_coverage.json')}、{link(OUT/'buffer_coverage_independent_review.json')}。这不表示网络或层本身不可实现；分块、流式、URAM等新实现需重新绑定和测量。因此P4平均耗时乘65×映射数仅能形成条件情景，缩减映射数量不能解决实现覆盖缺口。

## 交付与复现

- {link(OUT/'Experiment3_硬件测量试做.xlsx')}：D09、D07、{case_count}项状态、资源、{case_count*3}行板上范围记录、阶段耗时、原语盘点与来源绑定。
- {link(OUT/'D09_protocol.md')}：按草案执行本次pilot，实验5前正式冻结；逐项绑定计时语义、实际PL时钟与环境偏差证据。
- {link(OUT/'D07_mapping_recommendation.md')}：超过21天的条件性决策规则；保留低/高并行和必要的实测非支配折中点。
- {link(ROOT/'tools/hardware_measurement_pilot/README.md')}：重新生成与统计命令。

旧的HLS数据库、历史bitstream、JTAG可访问状态和以前的板上结果均不计入本次硬件通过数。

数据检查问题：{'; '.join(data['issues']) if data['issues'] else '未发现当前导入资料的身份绑定问题；这不等于尚缺硬件证据已通过。'}
"""
    previous_protocol={}
    protocol_path=OUT/'D09_protocol.md'
    if protocol_path.is_file():
        previous_sha=sha(protocol_path)
        archive_path=OUT/'protocol_revisions'/('D09_protocol_'+previous_sha[:16]+'.md')
        archive_path.parent.mkdir(exist_ok=True)
        if not archive_path.is_file():
            shutil.copy2(protocol_path,archive_path)
        if sha(archive_path)!=previous_sha:
            raise ValueError('D09 prior draft archive SHA256 mismatch')
        previous_protocol=dict(source=protocol_path.as_posix(),sha256=previous_sha,archive=archive_path.as_posix())
    (OUT / "report.md").write_text(report, encoding="utf-8")
    (OUT / "D09_protocol.md").write_text(f"# D09 测量协议草案\n\n版本：{PROTOCOL_VERSION}\n\n状态：DRAFT_EXECUTE_PILOT_FREEZE_BEFORE_EXPERIMENT5。草案可依据试做修订，正式实验5前冻结；原工作簿F:H保持空白，I列仍待冻结。先前D09文档按SHA256留存于protocol_revisions/；本次实际执行设置以逐run原始manifest、session和绑定的冻结runner为准，原summary与环境偏差不追改。\n\n本版差异：补充当前相位证据边界、实际CPU userspace条件、派生异常标签与有限条件排程；明确草案修订不等于正式冻结，也不把历史CPU performance偏差追认合规。\n\n正式冻结建议：若沿用当前板卡，将userspace列为明确允许的运行条件，冻结实际频率稳定性记录与偏差处理；若必须performance，则使用支持该模式的环境。继续保留3独立批次、每批50次预热与至少1000次正式调用、PL逐批及合并CV≤5%和host描述统计。冻结相位2为含局部控制开销的片上区间，完整动态集成RTL仿真仍记录NOT_RUN。改动以新版本重新采集正式数据，本次试做只作流程与参数选择依据。\n\n| 条目 | 执行要求 | 来源/证据 |\n|---|---|---|\n" + metadata_rows(D09) + "\n", encoding="utf-8")
    (OUT / "D07_mapping_recommendation.md").write_text("# D07 映射方案建议\n\n建议先2种不同物理配置起步，需要实测非支配折中时再增加第3种；六个查询策略可以复用物理行。完整65键可信日期仍NOT_ESTIMABLE，下列为P4实测mix条件情景，保留CPU userspace偏差，不属于正式成本表或D09条件资格。\n\n| 条目 | 决策规则 | 来源 |\n|---|---|---|\n" + metadata_rows(_d07_rows(data)) + "\n\n" + schedule_table + "\n", encoding="utf-8")
    (OUT / "extrapolation.json").write_text(json.dumps(data["extrapolation"], ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")
    (OUT / 'report_data.json').write_text(json.dumps(data,ensure_ascii=False,indent=2,allow_nan=False)+'\n',encoding='utf-8')
    validation = write_workbook(data, OUT / "Experiment3_硬件测量试做.xlsx")
    report_manifest = dict(generated_at=timestamp, attempt=attempt.as_posix(), validation=validation,
                           protocol_version=PROTOCOL_VERSION,previous_d09_document=previous_protocol,
                           native_pass_cases=native_pass, native_pass_fixtures=fixtures,
                           board_numeric_and_fixtures_pass_cases=runtime_pass,
                           measured_calls=measured_calls,warmup_calls=warmup_calls,independent_batches=batches,
                           cpu_condition_deviation_cases=cpu_deviation,
                           bound_phase_evidence_cases=sum(row['phase_evidence']['status']=='PASS_BOUND_PHASE_EVIDENCE' for row in cases),
                           annotation_case_count=data['annotations']['case_count'],annotation_total_rows=data['annotations']['total_rows'],
                           completed_automatic_driver_cases=data['automatic_driver_observations']['case_count'],
                           statuses={name: dict(Counter(r[name] for r in cases)) for name in ("native", "csim", "hls", "route", "bitstream", "board", "static_interface_acceptance", "board_execution", "board_correctness", "board_timing", "board_acceptance")},
                           source_sha256={path.as_posix(): sha(path) for path in
                                          (attempt/'source_manifest.json', attempt/'native_summary.json', attempt/'inventory.json', Path(__file__),
                                           Path(__file__).with_name('schedule_scenarios.py'), Path(__file__).with_name('anomaly_annotation.py'),
                                           Path(data['annotations']['source'])) if path.is_file()},
                           output_sha256={path.name:sha(path) for path in
                               (OUT/'report.md',OUT/'D09_protocol.md',OUT/'D07_mapping_recommendation.md',OUT/'extrapolation.json',
                                OUT/'report_data.json',OUT/'Experiment3_硬件测量试做.xlsx')},
                           issues=data["issues"], formal_cost_eligible=False)
    (OUT / "report_manifest.json").write_text(json.dumps(report_manifest, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")
    return report_manifest


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--attempt", type=Path, default=DEFAULT_ATTEMPT)
    parser.add_argument("--update-original", action="store_true", help="Apply the explicitly authorized original-workbook edits only")
    parser.add_argument("--stage-audit-only", action="store_true", help="Write a local-only stage audit without changing workbook/report.md")
    parser.add_argument("--successor", type=Path, default=OUT / "attempt05")
    args = parser.parse_args()
    if args.stage_audit_only:
        if args.update_original:
            parser.error("Stage audit must not update the original workbook")
        result = stage_audit(args.attempt, args.successor)
        print(json.dumps(dict(statuses=result["statuses"], build_only_observations={key: result["build_only_observations"][key]
            for key in ("completed_process_envelopes", "observed_peak_concurrent_envelopes", "observed_calendar_span_seconds", "summed_hls_route_process_seconds")}), ensure_ascii=False, indent=2))
        return
    if args.update_original:
        result = update_original_workbook()
        print(json.dumps({key:result[key] for key in ("status", "before_sha256", "after_sha256", "t21_row")}, ensure_ascii=False))
        return
    result = generate(args.attempt)
    print(json.dumps(dict(validation=result["validation"], statuses=result["statuses"]), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
